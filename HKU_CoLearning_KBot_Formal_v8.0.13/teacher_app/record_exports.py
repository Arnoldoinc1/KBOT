from __future__ import annotations

import io
import json
import re
from typing import Any, Iterable

from common.feedback_builder import build_record_review


EXPORT_HEADERS = [
    "Student",
    "Chapter",
    "Topic",
    "Question Type",
    "Question",
    "Student Answer",
    "Correct Answer",
    "AI Score",
    "Teacher Score",
    "AI Correct",
    "Teacher Verified",
    "AI Feedback",
    "Teacher Feedback",
    "Generation Source",
    "Grading Source",
    "Source Document",
    "Source Page",
    "Source Pages",
    "Created At",
    "Options",
    "Correct Option",
    "Matching Submitted Pairs",
    "Matching Correct Pairs",
    "Key Points",
    "Citations",
]


COLUMN_WIDTHS = {
    "Student": 16,
    "Chapter": 18,
    "Topic": 18,
    "Question Type": 16,
    "Question": 44,
    "Student Answer": 36,
    "Correct Answer": 36,
    "AI Score": 10,
    "Teacher Score": 12,
    "AI Correct": 12,
    "Teacher Verified": 16,
    "AI Feedback": 42,
    "Teacher Feedback": 42,
    "Generation Source": 18,
    "Grading Source": 18,
    "Source Document": 26,
    "Source Page": 11,
    "Source Pages": 18,
    "Created At": 22,
    "Options": 42,
    "Correct Option": 26,
    "Matching Submitted Pairs": 44,
    "Matching Correct Pairs": 44,
    "Key Points": 44,
    "Citations": 62,
}


def _text(value: object) -> str:
    if value is None:
        return ""
    return re.sub(r"\s+", " ", str(value)).strip()


def _multiline_text(value: object) -> str:
    if value is None:
        return ""
    text = str(value).replace("\r\n", "\n").replace("\r", "\n")
    lines = [re.sub(r"[ \t\f\v]+", " ", line).strip() for line in text.split("\n")]
    return "\n".join(line for line in lines if line)


def _decode_json(value: object) -> object:
    if not isinstance(value, str):
        return value
    raw = value.strip()
    if not raw:
        return ""
    try:
        return json.loads(raw)
    except (TypeError, json.JSONDecodeError):
        return raw


def _format_timestamp(value: object) -> str:
    if value is None:
        return ""
    formatter = getattr(value, "isoformat", None)
    if callable(formatter):
        try:
            return formatter(sep=" ", timespec="seconds")
        except (TypeError, ValueError):
            return _text(value)
    return _text(value)


def _format_options(review: dict[str, Any], payload: dict[str, Any]) -> str:
    options = review.get("options")
    if isinstance(options, list):
        lines: list[str] = []
        for option in options:
            if not isinstance(option, dict):
                continue
            label = _text(option.get("label"))
            option_text = _text(option.get("text"))
            if not label and not option_text:
                continue
            line = f"{label}. {option_text}" if label and option_text else option_text or label
            tags: list[str] = []
            if option.get("is_correct"):
                tags.append("CORRECT")
            if option.get("is_selected"):
                tags.append("STUDENT SELECTED")
            if tags:
                line += f" [{'; '.join(tags)}]"
            lines.append(line)
        if lines:
            return "\n".join(lines)

    left = payload.get("left")
    right = payload.get("right")
    if isinstance(left, list) or isinstance(right, list):
        left_values = [_text(item) for item in (left if isinstance(left, list) else []) if _text(item)]
        right_values = [_text(item) for item in (right if isinstance(right, list) else []) if _text(item)]
        lines = []
        if left_values:
            lines.append("Left: " + " | ".join(left_values))
        if right_values:
            lines.append("Right: " + " | ".join(right_values))
        return "\n".join(lines)
    return ""


def _format_pairs(value: object) -> str:
    if not isinstance(value, list):
        return ""
    lines = []
    for index, item in enumerate(value, start=1):
        text = _text(item)
        if text:
            lines.append(f"{index}. {text}")
    return "\n".join(lines)


def _format_key_points(payload: dict[str, Any]) -> str:
    points = payload.get("key_points")
    if isinstance(points, list):
        return _format_pairs(points)
    return _text(points)


def _format_citations(raw: object) -> str:
    citations = _decode_json(raw)
    if not isinstance(citations, list):
        return ""

    lines = []
    for index, item in enumerate(citations, start=1):
        if not isinstance(item, dict):
            continue
        source = _text(item.get("source_file") or item.get("filename") or item.get("source")) or "Unknown source"
        parts = [f"{index}. {source}"]
        section = _text(item.get("section") or item.get("chapter"))
        if section:
            parts.append(f"section: {section}")
        page = item.get("page")
        if page is not None and _text(page):
            parts.append(f"page: {page}")
        excerpt = _text(item.get("excerpt"))
        if len(excerpt) > 240:
            excerpt = excerpt[:237].rstrip() + "..."
        if excerpt:
            parts.append(f"excerpt: {excerpt}")
        lines.append(" | ".join(parts))
    return "\n".join(lines)


def _safe_int(value: object) -> int:
    try:
        return int(value or 0)
    except (TypeError, ValueError):
        return 0


def _safe_review(record: Any) -> dict[str, Any]:
    try:
        return build_record_review(record)
    except Exception:
        return {
            "question_type": _text(getattr(record, "question_type", "")) or "Unknown",
            "question": _text(getattr(record, "question", "")),
            "student_answer": _text(getattr(record, "answer", "")),
            "correct_answer": _text(getattr(record, "correct_answer", "")),
            "options": [],
            "submitted_pairs": [],
            "correct_pairs": [],
        }


def _record_row(record: Any) -> dict[str, Any]:
    review = _safe_review(record)
    payload = _decode_json(getattr(record, "options_or_pairs", ""))
    payload = payload if isinstance(payload, dict) else {}
    question_type = _text(review.get("question_type")) or _text(getattr(record, "question_type", ""))
    correct_answer = _text(review.get("correct_answer") or getattr(record, "correct_answer", ""))

    return {
        "Student": _text(getattr(record, "student_name", "")),
        "Chapter": _text(getattr(record, "chapter", "")),
        "Topic": _text(getattr(record, "topic", "")),
        "Question Type": question_type,
        "Question": _text(review.get("question") or getattr(record, "question", "")),
        "Student Answer": _text(review.get("student_answer") or getattr(record, "answer", "")),
        "Correct Answer": correct_answer,
        "AI Score": _safe_int(getattr(record, "score", 0)),
        "Teacher Score": getattr(record, "teacher_score", None) if getattr(record, "teacher_score", None) is not None else "",
        "AI Correct": "Correct" if bool(getattr(record, "is_correct", False)) else "Incorrect",
        "Teacher Verified": "Yes" if bool(getattr(record, "teacher_verified", False)) else "No",
        "AI Feedback": _multiline_text(getattr(record, "ai_feedback", "")),
        "Teacher Feedback": _multiline_text(getattr(record, "teacher_feedback", "")),
        "Generation Source": _text(getattr(record, "generation_source", "") or getattr(record, "source", "")),
        "Grading Source": _text(getattr(record, "grading_source", "")),
        "Source Document": _text(getattr(record, "source_document", "")),
        "Source Page": getattr(record, "source_page", None) if getattr(record, "source_page", None) is not None else "",
        "Source Pages": _text(getattr(record, "source_pages_json", "")),
        "Created At": _format_timestamp(getattr(record, "created_at", None)),
        "Options": _format_options(review, payload),
        "Correct Option": correct_answer if question_type in {"MCQ", "True/False"} else "",
        "Matching Submitted Pairs": _format_pairs(review.get("submitted_pairs")),
        "Matching Correct Pairs": _format_pairs(review.get("correct_pairs")),
        "Key Points": _format_key_points(payload),
        "Citations": _format_citations(getattr(record, "citations_json", "")),
    }


def build_records_xlsx(records: Iterable[Any]) -> tuple[io.BytesIO, int]:
    """Build a readable workbook from persisted student answer records."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Student Records"
    worksheet.freeze_panes = "A2"

    worksheet.append(EXPORT_HEADERS)
    rows = [_record_row(record) for record in records]
    for row in rows:
        worksheet.append([row.get(header, "") for header in EXPORT_HEADERS])

    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(color="FFFFFF", bold=True)
    header_alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    body_alignment = Alignment(vertical="top", wrap_text=True)
    border = Border(bottom=Side(style="thin", color="D9E2F3"))

    for cell in worksheet[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = header_alignment
        cell.border = border
    worksheet.row_dimensions[1].height = 28

    for row in worksheet.iter_rows(min_row=2, max_row=worksheet.max_row):
        for cell in row:
            cell.alignment = body_alignment
            cell.border = border

    for index, header in enumerate(EXPORT_HEADERS, start=1):
        worksheet.column_dimensions[get_column_letter(index)].width = COLUMN_WIDTHS.get(header, 20)
    worksheet.auto_filter.ref = f"A1:{get_column_letter(len(EXPORT_HEADERS))}{worksheet.max_row}"

    output = io.BytesIO()
    workbook.save(output)
    output.seek(0)
    return output, len(rows)
