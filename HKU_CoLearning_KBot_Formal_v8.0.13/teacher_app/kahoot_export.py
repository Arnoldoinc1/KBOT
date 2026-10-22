from __future__ import annotations

import io
import json
import re
from pathlib import Path
from typing import Any, Iterable


MCQ_TYPE_KEYS = {
    "mcq", "mcqtwo", "mcqfour", "mc", "multiplechoice",
    "multiplechoicequestion", "multiplechoicequestions",
}
TRUE_FALSE_TYPE_KEYS = {"truefalse", "tf", "boolean"}


def _type_key(value: object) -> str:
    return re.sub(r"[^a-z0-9]+", "", str(value or "").casefold())


def _safe_text(value: object, limit: int | None = None) -> str:
    text = str(value or "").encode("utf-8", errors="replace").decode("utf-8")
    text = re.sub(r"[\x00-\x08\x0B\x0C\x0E-\x1F]", " ", text)
    text = re.sub(r"\s+", " ", text).strip()
    return text[:limit] if limit is not None else text


def _decode(value: object) -> object:
    if not isinstance(value, str):
        return value
    text = value.strip()
    if not text:
        return ""
    try:
        return json.loads(text)
    except (json.JSONDecodeError, TypeError):
        return text


def extract_options(raw: object, limit: int | None = None) -> list[str]:
    decoded = _decode(raw)
    if isinstance(decoded, list):
        options = decoded
    elif isinstance(decoded, dict):
        options = decoded.get("options") or decoded.get("choices") or decoded.get("answers") or []
    else:
        options = []
    if not isinstance(options, list):
        return []
    cleaned = [_safe_text(option, limit) for option in options]
    return [item for item in cleaned if item][:4]


def extract_correct_hint(raw: object) -> str | None:
    decoded = _decode(raw)
    if not isinstance(decoded, dict):
        return None
    if "correct_index" in decoded:
        try:
            index = int(decoded["correct_index"])
        except (TypeError, ValueError):
            index = -1
        if 0 <= index <= 3:
            return str(index + 1)
    for key in ("correct_answer", "answer", "correct"):
        value = decoded.get(key)
        if value is not None and _safe_text(value):
            return _safe_text(value)
    return None


def parse_options_from_question(text: object) -> tuple[str, list[str]] | None:
    question = _safe_text(text)
    if not question:
        return None
    numeric = re.split(r"(?:^|\s)([1-4])[\.、\)]\s*", question)
    if len(numeric) >= 5:
        stem = numeric[0].strip()
        options = [numeric[index].strip() for index in range(2, len(numeric), 2) if numeric[index].strip()]
        if len(options) >= 2:
            return stem, options[:4]
    letters = re.split(r"(?:^|\s)\(?([A-Da-d])[\)\.、]\s*", question)
    if len(letters) >= 5:
        stem = letters[0].strip()
        options = [letters[index].strip() for index in range(2, len(letters), 2) if letters[index].strip()]
        if len(options) >= 2:
            return stem, options[:4]
    return None


def extract_options_from_feedback(raw: object, limit: int | None = None) -> list[str]:
    text = str(raw or "")
    if not text:
        return []
    options: dict[str, str] = {}
    for line in text.splitlines():
        match = re.match(r"^\s*([A-Da-d])[\.\)]\s+(.+?)\s*(?:\[(?:CORRECT|STUDENT SELECTED|CORRECT\s*;\s*STUDENT SELECTED|STUDENT SELECTED\s*;\s*CORRECT)\])?\s*$", line)
        if not match:
            continue
        option = re.sub(r"\s*\[(?:CORRECT|STUDENT SELECTED)(?:\s*;\s*(?:CORRECT|STUDENT SELECTED))*\]\s*$", "", match.group(2), flags=re.IGNORECASE)
        options[match.group(1).upper()] = _safe_text(option, limit)
    ordered = [options[letter] for letter in "ABCD" if options.get(letter)]
    return ordered[:4]


def extract_correct_from_feedback(raw: object) -> str | None:
    text = str(raw or "")
    for line in text.splitlines():
        match = re.match(r"^\s*Correct option:\s*([A-Da-d]|[1-4])(?:[\.\):\-]\s*(.*))?$", line, re.IGNORECASE)
        if match:
            return match.group(1).upper()
    for line in text.splitlines():
        match = re.match(r"^\s*([A-Da-d])[\.\)]\s+.+?\s*\[(?:CORRECT|CORRECT\s*;\s*STUDENT SELECTED|STUDENT SELECTED\s*;\s*CORRECT)\]\s*$", line, re.IGNORECASE)
        if match:
            return match.group(1).upper()
    return None


def _normalized_text(value: object) -> str:
    return re.sub(r"\s+", " ", _safe_text(value).casefold()).strip()


def _question_fingerprint(kind: str, stem: str, options: list[str], correct: str) -> str:
    """Create a stable content key for duplicate Kahoot questions."""
    return json.dumps(
        [
            kind,
            _normalized_text(stem),
            [_normalized_text(option) for option in options],
            _normalized_text(correct),
        ],
        ensure_ascii=False,
        separators=(",", ":"),
    )


def resolve_correct_answer(raw_answer: object, options: list[str]) -> str | None:
    decoded = _decode(raw_answer)
    if isinstance(decoded, list):
        answer = ",".join(str(item) for item in decoded)
    elif isinstance(decoded, (int, float)):
        answer = str(int(decoded))
    else:
        answer = _safe_text(decoded)
    if not answer or not options:
        return None

    target = _normalized_text(answer)
    for index, option in enumerate(options, start=1):
        if target == _normalized_text(option):
            return str(index)

    match = re.match(r"^\s*([A-Da-d]|[1-4])[\)\.\:\-]\s*", answer)
    if match:
        token = match.group(1).upper()
        return str(ord(token) - ord("A") + 1) if token in {"A", "B", "C", "D"} else token
    if re.fullmatch(r"[A-Da-d]", answer):
        return str(ord(answer.upper()) - ord("A") + 1)
    if re.fullmatch(r"[1-4]", answer):
        return answer
    if answer == "0":
        return "1"
    return None


def eligible_questions(records: Iterable[Any]) -> list[dict[str, Any]]:
    questions: dict[str, dict[str, Any]] = {}
    for record in records:
        kind = _type_key(getattr(record, "question_type", ""))
        if kind not in MCQ_TYPE_KEYS and kind not in TRUE_FALSE_TYPE_KEYS:
            continue
        stem = _safe_text(getattr(record, "question", ""), 120)
        if not stem:
            continue
        feedback = getattr(record, "ai_feedback", "")
        if kind in TRUE_FALSE_TYPE_KEYS:
            raw_options = ["True", "False"]
        else:
            raw_options = extract_options(getattr(record, "options_or_pairs", ""), limit=None)
            if len(raw_options) < 2:
                raw_options = extract_options_from_feedback(feedback, limit=None)
            if len(raw_options) < 2:
                parsed = parse_options_from_question(stem)
                if parsed:
                    stem, raw_options = parsed
        if len(raw_options) < 2:
            continue

        correct = None
        correct_candidates = (
            getattr(record, "correct_answer", ""),
            extract_correct_hint(getattr(record, "options_or_pairs", "")),
            extract_correct_from_feedback(feedback),
        )
        for candidate in correct_candidates:
            correct = resolve_correct_answer(candidate, raw_options)
            if correct is not None:
                break
        if correct is None:
            continue

        # Kahoot limits answer text to 75 characters. Resolve the correct index
        # against the complete stored options first, then truncate only for XLSX.
        # Deduplicate by the complete quiz content, not by the stem alone.
        # This keeps identical student attempts together while preserving
        # same-stem questions that have different options or correct answers.
        key = _question_fingerprint(kind, stem, raw_options, correct)
        options = ([_safe_text(option, 75) for option in raw_options] + ["", "", "", ""])[:4]
        questions.setdefault(key, {"stem": stem, "options": options, "correct": correct})
    return list(questions.values())


def build_kahoot_xlsx(records: Iterable[Any], template_path: str | Path) -> tuple[io.BytesIO, int]:
    from openpyxl import Workbook, load_workbook

    template = Path(template_path)
    if template.exists():
        workbook = load_workbook(template)
        worksheet = workbook.active
        for row in worksheet.iter_rows(min_row=9, max_row=max(9, worksheet.max_row)):
            for cell in row:
                cell.value = None
    else:
        workbook = Workbook()
        worksheet = workbook.active
        worksheet.title = "Quiz"
        worksheet.cell(row=8, column=1, value="Question")
        for index in range(4):
            worksheet.cell(row=8, column=2 + index, value=f"Answer {index + 1}")
        worksheet.cell(row=8, column=6, value="Time limit")
        worksheet.cell(row=8, column=7, value="Correct answer(s)")

    questions = eligible_questions(records)
    for row_number, item in enumerate(questions, start=9):
        worksheet.cell(row=row_number, column=1, value=item["stem"])
        for column_number, option in enumerate(item["options"], start=2):
            worksheet.cell(row=row_number, column=column_number, value=_safe_text(option, 75))
        worksheet.cell(row=row_number, column=6, value=60)
        worksheet.cell(row=row_number, column=7, value=item["correct"])

    output = io.BytesIO()
    workbook.save(output)
    output.seek(0)
    return output, len(questions)
