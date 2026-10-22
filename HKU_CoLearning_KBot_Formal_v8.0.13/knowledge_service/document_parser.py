from __future__ import annotations

import io
import re
from dataclasses import dataclass
from pathlib import Path
from typing import Any

from docx import Document
from pypdf import PdfReader


@dataclass(frozen=True)
class ParsedBlock:
    text: str
    kind: str = "paragraph"


@dataclass(frozen=True)
class ParsedPage:
    page: int
    text: str
    chapter: str
    section: str
    blocks: tuple[ParsedBlock, ...] = ()


BOILERPLATE = (
    "most materials are from the text",
    "for use in lectures only",
    "not for duplication",
    "courtesy:",
)
MAJOR_HEADING = re.compile(
    r"^\s*(?:chapter|unit|module|lecture)\s+(?!contents?\b)([A-Za-z0-9][\w.-]*)\b\s*[:\-]?\s*(.*)$",
    re.IGNORECASE,
)
NUMBERED_HEADING = re.compile(r"^\d+(?:\.\d+){0,3}\s+[A-Z][^.!?]{2,100}$")
BULLET_PREFIX = re.compile(r"^(?:[\u2022\u25e6\uf06e*\-]|\u2013|\u2014|\d+(?:\.\d+)*[.)])\s*")
NUMBERED_BULLET = re.compile(r"^[\u2022\u25e6\uf06e*\-]?\s*\d+(?:\.\d+)*[.)]\s+")
DEFINITION_PREFIX = re.compile(r"^[A-Za-z][A-Za-z0-9/().,\- ]{0,80}:\s+")
OPEN_ENDING = re.compile(
    r"\b(?:and|or|to|of|for|with|from|into|through|including|such as|due to)\s*$",
    re.IGNORECASE,
)
NOISE_PATTERNS = (
    re.compile(r"^https?://", re.IGNORECASE),
    re.compile(r"^\d+\s*[A-Z][A-Za-z].{80,}"),
    re.compile(r"\b(?:doi\.org|isbn|international symposium on electronic imaging)\b", re.IGNORECASE),
)


def clean_text(value: object) -> str:
    text = str(value or "").encode("utf-8", errors="replace").decode("utf-8").replace("\x00", " ").replace("\u00ad", "")
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def _strip_inline_noise(value: object) -> str:
    text = str(value or "")
    text = re.sub(r"https?://\S+", "", text, flags=re.IGNORECASE)
    return text


def clean_lines(value: object) -> list[str]:
    lines: list[str] = []
    for raw in str(value or "").splitlines():
        line = clean_text(_strip_inline_noise(raw))
        if not line or any(marker in line.casefold() for marker in BOILERPLATE):
            continue
        if re.fullmatch(r"\d{1,4}", line):
            continue
        if any(pattern.search(line) for pattern in NOISE_PATTERNS):
            continue
        lines.append(line)
    return lines


def infer_document_label(filename: str) -> str:
    stem = re.sub(r"^[0-9a-f]{16,}__", "", Path(filename).stem, flags=re.IGNORECASE)
    return re.sub(r"[_-]+", " ", stem).strip() or "Course material"


def _plausible_title(line: str) -> bool:
    if len(line) < 3 or len(line) > 110 or line.startswith(("\u2022", "-", "\u2013", "\u2014")):
        return False
    if line.lower().startswith(("http://", "https://")) or line.endswith((".", ";")):
        return False
    words = line.split()
    if len(words) > 14:
        return False
    if "=" in line or line in {"?", "x", "X"}:
        return False
    return True


def page_title(text: str, fallback: str) -> str:
    lines = clean_lines(text)
    candidates = [(index, line) for index, line in enumerate(lines[:24]) if _plausible_title(line)]
    if not candidates:
        return fallback

    def score(candidate: tuple[int, str]) -> float:
        index, line = candidate
        words = line.split()
        value = 4.0 if index == 0 else max(0.0, 1.5 - index * 0.08)
        value += 6.0 if MAJOR_HEADING.search(line) or NUMBERED_HEADING.match(line) or ":" in line else 0.0
        value += 3.0 if sum(word[:1].isupper() for word in words) >= max(1, len(words) * 0.6) else 0.0
        value += 2.0 if line.isupper() else 0.0
        value += 3.0 if re.search(r"\b(?:GPS|system|segment|component|principle|background|error|methodolog|technology|network)\b", line, re.IGNORECASE) else 0.0
        value -= 4.0 if re.search(r"\b(?:is|are|was|were|would|could|then)\b", line, re.IGNORECASE) else 0.0
        return value

    title = lines[0] if _plausible_title(lines[0]) else max(candidates, key=score)[1]
    if len(lines) > 1 and title == lines[0] and lines[1].isupper() and 2 <= len(lines[1]) <= 12 and lines[1] not in title:
        title = f"{title} ({lines[1]})"
    return title[:160]


def _normalize_line(value: object) -> str:
    text = re.sub(r"\s+", " ", _strip_inline_noise(value).strip())
    text = re.sub(r"^([•\u2022\uf06e\u25e6]\s*[A-Za-z][A-Za-z0-9/()]{0,30})-(?=\s)", r"\1 -", text)
    return text


def _is_noise_line(line: str) -> bool:
    text = _normalize_line(line)
    if not text:
        return True
    lowered = text.casefold()
    if any(marker in lowered for marker in BOILERPLATE):
        return True
    if lowered in {"vertical horizontal"}:
        return True
    if re.fullmatch(r"\d{1,4}", text):
        return True
    if re.fullmatch(r"[|\-_=]{3,}", text):
        return True
    return any(pattern.search(text) for pattern in NOISE_PATTERNS)


def _block_kind(text: str) -> str:
    first = text.splitlines()[0].strip() if text else ""
    if BULLET_PREFIX.match(first):
        return "bullet"
    if DEFINITION_PREFIX.match(first):
        return "definition"
    if re.match(r"^(?:example|figure|table|q:)\b", first, re.IGNORECASE):
        return "heading"
    if _plausible_title(first) and not first.endswith((".", "?", "!")):
        words = first.split()
        capitalized_words = sum(word[:1].isupper() for word in words)
        if (
            first[:1].isupper()
            and (len(words) <= 3 or first.isupper())
            or capitalized_words >= 2
            or NUMBERED_HEADING.match(first)
            or MAJOR_HEADING.search(first)
        ):
            return "heading"
        if ":" in first and capitalized_words >= 2:
            return "heading"
    return "paragraph"


def _is_section_heading(block: ParsedBlock) -> bool:
    if block.kind != "heading":
        return False
    text = block.text.strip()
    if re.match(r"^(?:q:|example|figure|table)\b", text, re.IGNORECASE):
        return False
    if text.casefold().startswith(("feature ", "column ")):
        return False
    if text.endswith((",", ";")):
        return False
    return _plausible_title(text)


def _line_starts_new_block(
    *,
    current_lines: list[str],
    current_kind: str,
    start_indent: int,
    text: str,
    indent: int,
) -> bool:
    if not current_lines:
        return False
    kind = _block_kind(text)
    previous = current_lines[-1].strip()

    if current_kind == "heading":
        return True
    if current_kind in {"bullet", "definition"}:
        if kind in {"bullet", "definition"}:
            return True
        if kind == "heading":
            # PDF line wrapping can make a continuation line look like a
            # heading (for example "1/683 W ..." or a capitalized proper
            # name).  Keep such lines inside the current bullet unless they
            # are unmistakably a real structural heading.
            return bool(MAJOR_HEADING.match(text) or NUMBERED_HEADING.match(text) or text.isupper())
        return False
    if kind in {"heading", "definition"}:
        return True
    if indent > start_indent + 1:
        return False
    if previous.endswith((':', '-', "\u2013", "\u2014", "/", "(", ",")) or OPEN_ENDING.search(previous):
        return False
    if text[:1].islower():
        return False
    return len(previous) > 80 or previous.endswith((".", "?", "!"))


def _text_blocks(text: str) -> list[ParsedBlock]:
    lines: list[tuple[int, str]] = []
    for raw in text.splitlines():
        if _is_noise_line(raw):
            continue
        indent = len(raw) - len(raw.lstrip(" "))
        line = _normalize_line(raw)
        if line:
            lines.append((indent, line))

    blocks: list[ParsedBlock] = []
    current: list[str] = []
    current_kind = ""
    start_indent = 0
    for indent, line in lines:
        if not current:
            current = [line]
            current_kind = _block_kind(line)
            start_indent = indent
            continue
        if _line_starts_new_block(
            current_lines=current,
            current_kind=current_kind,
            start_indent=start_indent,
            text=line,
            indent=indent,
        ):
            block_text = clean_text("\n".join(current))
            blocks.append(ParsedBlock(block_text, _block_kind(block_text)))
            current = [line]
            current_kind = _block_kind(line)
            start_indent = indent
        else:
            current.append(line)
    if current:
        block_text = clean_text("\n".join(current))
        blocks.append(ParsedBlock(block_text, _block_kind(block_text)))
    return blocks


def _table_row_blocks(table: Any) -> list[ParsedBlock]:
    if not table or len(table) < 2:
        return []
    header = [clean_text(cell).replace("\n", " ") if cell else "" for cell in table[0]]
    if not any(header):
        return []

    blocks: list[ParsedBlock] = []
    for row in table[1:]:
        cells = [clean_text(cell).replace("\n", " ") if cell else "" for cell in row]
        if not any(cells):
            continue
        label = cells[0] or header[0] or "Table entry"
        parts: list[str] = []
        for index, cell in enumerate(cells[1:], start=1):
            if not cell:
                continue
            column = header[index] if index < len(header) and header[index] else f"Column {index + 1}"
            parts.append(f"{column}: {cell}.")
        if parts:
            blocks.append(ParsedBlock(clean_text(f"{label}. {' '.join(parts)}"), "table_row"))
    return blocks


# Older PDF exports (including HKU's original Bayes slides) encode SymbolMT
# letters and operators as U+F0xx private-use characters. They look correct in
# the PDF's embedded font but are unreadable to RAG and quiz validators. Decode
# the glyph *only when the font is Symbol*, otherwise the same code point may
# mean a Wingdings pictogram. Do not infer missing fractions or equation order.
_PDF_SYMBOL_MAP = {34: '∀', 40: '(', 41: ')', 43: '+', 45: '−', 60: '<', 61: '=', 62: '>', 91: '[', 93: ']', 97: 'α', 99: 'χ', 108: 'λ', 109: 'μ', 112: 'π', 115: 'σ', 119: 'ω', 163: '≤', 165: '∞', 179: '≥', 180: '×', 185: '≠', 206: '∈', 215: '⋅', 229: '∑', 242: '∫'}


def _decoded_pdf_chars(chars: list[dict]) -> list[dict]:
    output: list[dict] = []
    for char in chars:
        value = str(char.get("text") or "")
        font = str(char.get("fontname") or "").split("+")[-1].casefold()
        replacement = value
        if len(value) == 1 and 0xF000 <= ord(value) <= 0xF0FF:
            if "symbol" in font:
                replacement = _PDF_SYMBOL_MAP.get(ord(value) - 0xF000, value)
            elif "wingdings" in font and ord(value) == 0xF06E:
                replacement = "•"
        output.append({**char, "text": replacement} if replacement != value else char)
    return output


def _extract_pdf_pages(data: bytes) -> list[tuple[int, str, tuple[ParsedBlock, ...]]]:
    try:
        import pdfplumber
    except ImportError:
        reader = PdfReader(io.BytesIO(data))
        extracted = [
            (index + 1, clean_text(page.extract_text() or ""), ())
            for index, page in enumerate(reader.pages)
        ]
        if any(
            any(0xE000 <= ord(char) <= 0xF8FF for char in page_text)
            for _page, page_text, _blocks in extracted
        ):
            raise RuntimeError(
                "This legacy PDF contains Symbol-font equations that pypdf cannot "
                "extract reliably. Install pdfplumber with 'python -m pip install "
                "-r requirements.txt' and then upload the PDF again."
            )
        return extracted

    pages: list[tuple[int, str, tuple[ParsedBlock, ...]]] = []
    with pdfplumber.open(io.BytesIO(data)) as document:
        for index, page in enumerate(document.pages, start=1):
            # Extract through font-aware characters, not plain extract_text(),
            # which loses the fact that U+F077 came from the Symbol omega glyph.
            # Most PDFs contain no private-use Symbol glyphs and use the old
            # high-performance path unchanged.
            needs_mapping = any(
                len(char.get("text", "")) == 1
                and 0xF000 <= ord(char["text"]) <= 0xF0FF
                and any(marker in str(char.get("fontname") or "").casefold()
                        for marker in ("symbol", "wingdings"))
                for char in page.chars
            )
            if needs_mapping:
                from pdfplumber.utils import extract_text as extract_char_text
                raw = extract_char_text(_decoded_pdf_chars(page.chars), layout=True)
            else:
                raw = page.extract_text(layout=True) or page.extract_text() or ""
            text = clean_text(raw)
            blocks = _text_blocks(raw)
            table_blocks: list[ParsedBlock] = []
            try:
                for table in page.extract_tables() or []:
                    table_blocks.extend(_table_row_blocks(table))
            except Exception:
                table_blocks = []
            if table_blocks:
                heading_blocks = [block for block in blocks[:3] if _is_section_heading(block)]
                blocks = [*heading_blocks, *table_blocks]
                text = clean_text("\n".join(block.text for block in blocks))
            pages.append((index, text, tuple(blocks)))
    return pages


def _extract_docx_pages(data: bytes) -> list[tuple[int, str, tuple[ParsedBlock, ...]]]:
    document = Document(io.BytesIO(data))
    blocks: list[ParsedBlock] = []
    for paragraph in document.paragraphs:
        text = clean_text(paragraph.text)
        if text:
            blocks.append(ParsedBlock(text, _block_kind(text)))
    for table in document.tables:
        rows = [[cell.text for cell in row.cells] for row in table.rows]
        blocks.extend(_table_row_blocks(rows))
    body = clean_text("\n".join(block.text for block in blocks))
    return [(1, body, tuple(blocks))] if body else []


def _extract_pptx_pages(data: bytes) -> list[tuple[int, str, tuple[ParsedBlock, ...]]]:
    try:
        from pptx import Presentation
    except Exception as exc:
        raise ValueError("PPTX course materials require python-pptx to be installed.") from exc

    presentation = Presentation(io.BytesIO(data))
    pages: list[tuple[int, str, tuple[ParsedBlock, ...]]] = []
    for index, slide in enumerate(presentation.slides, start=1):
        blocks: list[ParsedBlock] = []
        for shape in slide.shapes:
            if getattr(shape, "has_table", False):
                rows = [[cell.text for cell in row.cells] for row in shape.table.rows]
                blocks.extend(_table_row_blocks(rows))
                continue
            if not getattr(shape, "has_text_frame", False):
                continue
            for paragraph in shape.text_frame.paragraphs:
                text = clean_text("".join(run.text for run in paragraph.runs) or paragraph.text)
                if text:
                    blocks.append(ParsedBlock(text, _block_kind(text)))
        body = clean_text("\n".join(block.text for block in blocks))
        if body:
            pages.append((index, body, tuple(blocks)))
    return pages


def _raw_pages(data: bytes, filename: str) -> list[tuple[int, str, tuple[ParsedBlock, ...]]]:
    suffix = Path(filename).suffix.casefold()
    if suffix == ".pdf":
        return _extract_pdf_pages(data)
    if suffix == ".docx":
        return _extract_docx_pages(data)
    if suffix == ".pptx":
        return _extract_pptx_pages(data)
    if suffix == ".txt":
        body = clean_text(data.decode("utf-8", errors="replace"))
        return [(1, body, tuple(_text_blocks(body)))] if body else []
    raise ValueError("Only PDF, PPTX, DOCX, and TXT course materials are supported.")


def parse_document(data: bytes, filename: str) -> list[ParsedPage]:
    default = infer_document_label(filename)
    current_chapter = default
    output: list[ParsedPage] = []
    for page_number, raw, blocks in _raw_pages(data, filename):
        lines = clean_lines(raw)
        text = clean_text("\n".join(lines))
        if not text:
            continue
        for line in lines[:3]:
            match = MAJOR_HEADING.search(line)
            if match:
                suffix = clean_text(match.group(2))
                current_chapter = f"Chapter {match.group(1)}" + (f": {suffix}" if suffix else "")
                break
        output.append(ParsedPage(page_number, text, current_chapter, page_title(text, current_chapter), blocks))
    return output


def _fallback_blocks(text: str) -> list[ParsedBlock]:
    blocks = _text_blocks(text)
    if len(blocks) == 1 and blocks[0].text == clean_text(text):
        parts = [clean_text(part) for part in re.split(r"(?=\n?[\u2022\uf06e\u2013\-]\s+)", text) if clean_text(part)]
        if len(parts) > 1:
            return [ParsedBlock(part, _block_kind(part)) for part in parts]
    return blocks or [ParsedBlock(clean_text(text), "paragraph")]


def _chunk_segments(blocks: list[ParsedBlock], chunk_size: int, soft_size: int) -> list[str]:
    chunks: list[str] = []
    buffer: list[str] = []
    buffer_size = 0
    active_heading = ""

    def flush() -> None:
        nonlocal buffer, buffer_size
        if buffer:
            chunks.append(clean_text("\n".join(buffer)))
            buffer = []
            buffer_size = 0

    for block in blocks:
        text = clean_text(block.text)
        if not text:
            continue
        if _is_section_heading(block):
            active_heading = text
            if buffer_size >= 120:
                flush()
            if not buffer:
                buffer = [text]
                buffer_size = len(text)
            continue

        piece = text
        if active_heading and not buffer:
            buffer = [active_heading]
            buffer_size = len(active_heading)
        projected = buffer_size + len(piece) + (1 if buffer else 0)
        if buffer and (
            projected > chunk_size
            or block.kind in {"definition", "table_row"} and buffer_size >= soft_size
            or NUMBERED_BULLET.match(piece) and buffer_size >= soft_size
        ):
            flush()
            if active_heading and active_heading != piece:
                buffer = [active_heading]
                buffer_size = len(active_heading)
        if len(piece) > chunk_size:
            flush()
            start = 0
            while start < len(piece):
                chunks.append(clean_text(piece[start : start + chunk_size]))
                start += max(1, chunk_size)
            continue
        buffer.append(piece)
        buffer_size += len(piece) + (1 if buffer_size else 0)
    flush()
    return chunks


def split_chunks(pages: list[ParsedPage], chunk_size: int = 900, overlap: int = 0) -> list[dict]:
    output: list[dict] = []
    soft_size = max(260, min(650, int(chunk_size * 0.72)))
    for page in pages:
        blocks = list(page.blocks) if page.blocks else _fallback_blocks(page.text)
        chunks = _chunk_segments(blocks, max(350, chunk_size), soft_size)
        if not chunks:
            chunks = [page.text]
        for chunk in chunks:
            if len(chunk) < 30:
                continue
            output.append({"page": page.page, "chapter": page.chapter, "section": page.section, "text": chunk})
    return output


def fact_sentences(text: str) -> list[str]:
    normalized = re.sub(r"\n[\u2022\uf06e]\s*", ". ", text)
    normalized = re.sub(r"\n[\u2013\-]\s*", " ", normalized)
    normalized = re.sub(r"\n+", " ", normalized)
    facts: list[str] = []
    for sentence in re.split(r"(?<=[.!?])\s+", normalized):
        sentence = clean_text(sentence).strip("\u2022\uf06e\u2013- ")
        words = re.findall(r"[A-Za-z][A-Za-z0-9'/-]*", sentence)
        lowered = sentence.casefold()
        bibliography = re.search(r"\b(?:theory,? practice and applications|further reading|references?|copyright|isbn)\b", lowered)
        if 5 <= len(words) <= 65 and 25 <= len(sentence) <= 420 and "http://" not in lowered and "https://" not in lowered and not bibliography:
            facts.append(sentence)
    return list(dict.fromkeys(facts))


def _valid_term(value: str) -> bool:
    term = clean_text(value).strip("\u2022\uf06e\u2013-?,:;. ")
    words = term.split()
    blocked = {
        "introduction", "overview", "summary", "example", "applications", "how", "what",
        "the", "this", "that", "there", "for", "it", "a", "an", "to", "use", "using",
        "class", "discussion", "discussions",
    }
    if not (1 <= len(words) <= 7 and 2 <= len(term) <= 70):
        return False
    if words[0].casefold() in blocked or any(word.casefold() in {"that", "which", "who"} for word in words):
        return False
    if term[:1].islower() and not term.isupper():
        return False
    if term.endswith((".", ":", "?")) or not all(term.count(left) == term.count(right) for left, right in (("(", ")"), ("[", "]"))):
        return False
    return not re.search(r"\b(?:consists? of|components? include)\b", term, re.IGNORECASE)


def explicit_term_pairs(pages: list[ParsedPage]) -> list[dict]:
    pairs: list[dict] = []
    title_pages = [(page.section.casefold(), page) for page in pages]
    for page in pages:
        for fact in fact_sentences(page.text):
            match = re.match(r"^(.{2,70}?)\s+(?:is|are|means|refers to|describes)\s+(.{12,330})$", fact, re.IGNORECASE)
            if match and _valid_term(match.group(1)):
                pairs.append({"term": clean_text(match.group(1)), "definition": clean_text(match.group(2)), "category": page.section, "page": page.page})
            match = re.match(r"^(.{2,70}?):\s+(.{12,330})$", fact)
            if match and _valid_term(match.group(1)):
                pairs.append({"term": clean_text(match.group(1)), "definition": clean_text(match.group(2)), "category": page.section, "page": page.page})

        # Resolve explicit component lists against later slide titles and facts.
        if re.search(r"\b(?:consists? of|components? include|major components?)\b", page.text, re.IGNORECASE):
            lines = clean_lines(page.text)
            terms = [line.strip("\u2022\uf06e\u2013-?,:;. ") for line in lines[1:] if _valid_term(line.strip("\u2022\uf06e\u2013-?,:;. "))]
            for term in terms[:10]:
                term_tokens = set(re.findall(r"[a-z0-9]+", term.casefold()))
                ranked_targets = []
                for title, candidate in title_pages:
                    title_tokens = set(re.findall(r"[a-z0-9]+", title))
                    similarity = len(term_tokens & title_tokens) / len(term_tokens | title_tokens) if term_tokens and title_tokens else 0.0
                    if term.casefold() == title:
                        similarity += 1.0
                    ranked_targets.append((similarity, -abs(candidate.page - page.page), candidate))
                best = max(ranked_targets, default=(0.0, 0, None), key=lambda item: (item[0], item[1]))
                target = best[2] if best[0] >= 0.45 else None
                if not target or target.page == page.page:
                    continue
                definitions = fact_sentences(target.text)
                definition = next((item for item in definitions if term.casefold() not in item.casefold() or len(item) > len(term) + 25), "")
                if definition:
                    pairs.append({"term": term, "definition": definition, "category": page.section, "page": target.page})
    unique: list[dict] = []
    seen: set[tuple[str, str]] = set()
    for pair in pairs:
        key = (pair["term"].casefold(), pair["definition"].casefold())
        if key not in seen:
            seen.add(key)
            unique.append(pair)
    return unique
