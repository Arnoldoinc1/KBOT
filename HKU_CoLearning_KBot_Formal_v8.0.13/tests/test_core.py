from __future__ import annotations

import hashlib
import json
import os
import shutil
import subprocess
import tempfile
import unittest
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import patch

from sqlalchemy import create_engine, text
from sqlalchemy.orm import sessionmaker
from jinja2 import Environment, FileSystemLoader

from common.coze_transport import AgentConfig, AgentKind, CozeAgentRouter, CozeTransportError
from common.config import app_secret, database_url, env_bool, service_secret
from common.grading_engine import grade_objective, grade_short_answer, local_semantic_grade, normalize_grading
from common.json_recovery import parse_json_like
from common.notes_engine import NOT_FOUND, answer_notes, local_notes_answer, normalize_notes
from common.quiz_engine import FocusFact, QuizValidationError, generate_quiz, local_quiz, validate_grounding
from common.quiz_flow import request_matches_current, skip_transition, transition
from common.quiz_parser import normalize_quiz
from common.schema_compat import ensure_legacy_schema
from knowledge_service.repository import KnowledgeRepository
from knowledge_service.retrieval import lexical_score
from teacher_app.kahoot_export import build_kahoot_xlsx, eligible_questions
from teacher_app.record_search import hybrid_search_records


ROOT = Path(__file__).resolve().parents[1]
REAL_PDF = ROOT / "test_data" / "Ch-3_Enabling_Technologies_2025.pdf"


class DisabledRouter:
    def config(self, kind):
        return AgentConfig(AgentKind(kind), "https://api.coze.test", "", "")


class FailingRouter:
    def config(self, kind):
        return AgentConfig(AgentKind(kind), "https://api.coze.test", "token", f"{AgentKind(kind).value}-bot")

    def request(self, *_args, **_kwargs):
        raise CozeTransportError("simulated timeout or rate limit")


class StaticRouter:
    def __init__(self, active_kind, content):
        self.active_kind = AgentKind(active_kind)
        self.content = content

    def config(self, kind):
        active = AgentKind(kind) == self.active_kind
        return AgentConfig(AgentKind(kind), "https://api.coze.test", "token" if active else "", "bot" if active else "")

    def request(self, kind, _prompt, _user):
        if AgentKind(kind) != self.active_kind:
            raise AssertionError("wrong Agent route")
        return type("StaticResponse", (), {"content": self.content})()


class JsonRecoveryTests(unittest.TestCase):
    def test_complete_fenced_prefixed_and_truncated(self):
        values = [
            '{"type":"mcq","question":"Q"}',
            '```json\n{"type":"mcq","question":"Q"}\n```',
            'Here is the result: {"type":"mcq","question":"Q"} done',
            '{"type":"matching","p":[["A","Definition A"],["B","Definition B"]',
        ]
        for raw in values:
            with self.subTest(raw=raw[:30]):
                self.assertIsInstance(parse_json_like(raw), dict)


class ConfigurationContractTests(unittest.TestCase):
    def test_mysql_url_is_preserved_and_required_env_names_exist(self):
        previous = os.environ.get("DATABASE_URL")
        value = "mysql+pymysql://kbot:password@db:3306/edu_qa?charset=utf8mb4"
        try:
            os.environ["DATABASE_URL"] = value
            self.assertEqual(database_url(), value)
        finally:
            if previous is None:
                os.environ.pop("DATABASE_URL", None)
            else:
                os.environ["DATABASE_URL"] = previous
        env_text = (ROOT / ".env.example").read_text(encoding="utf-8")
        required = {"APP_SECRET_KEY", "DATABASE_URL", "KNOWLEDGE_SERVICE_URL", "KNOWLEDGE_ADMIN_KEY", "KNOWLEDGE_READ_KEY", "KNOWLEDGE_DATA_DIR", "COZE_API_BASE", "COZE_TOKEN", "COZE_QUIZ_BOT_ID", "COZE_GRADING_BOT_ID", "COZE_NOTES_BOT_ID", "COZE_TIMEOUT", "COZE_STREAM_ENABLED", "TEACHER_HOST", "TEACHER_PORT", "STUDENT_HOST", "STUDENT_PORT", "KNOWLEDGE_HOST", "KNOWLEDGE_PORT"}
        present = {line.split("=", 1)[0].strip() for line in env_text.splitlines() if "=" in line and not line.lstrip().startswith("#")}
        self.assertFalse(required - present)

    def test_mysql_initialization_has_all_shared_tables(self):
        sql = (ROOT / "migrations" / "mysql_init.sql").read_text(encoding="utf-8").casefold()
        for table in ("users", "student_qa_records", "notes_qa_records", "quiz_sessions"):
            self.assertIn(f"create table if not exists {table}", sql)
        self.assertIn("generation_failures_json", sql)
        self.assertIn("skipped_types_json", sql)


class QuizParserTests(unittest.TestCase):
    def test_matching_aliases_and_mappings(self):
        payloads = [
            {"p": [["A", "Definition A"], ["B", "Definition B"]]},
            {"pairs": [{"term": "A", "definition": "Definition A"}, {"term": "B", "definition": "Definition B"}]},
            {"matching_pairs": [{"left": "A", "right": "Definition A"}, {"left": "B", "right": "Definition B"}]},
            {"left": ["A", "B"], "right": ["Definition A", "Definition B"]},
            {"l": ["A", "B"], "r": ["Definition B", "Definition A"], "mapping": [2, 1]},
            {"terms": ["A", "B"], "definitions": ["Definition B", "Definition A"], "correct_mapping": [1, 0]},
            {"pairs": {"A": "Definition A", "B": "Definition B"}},
        ]
        for payload in payloads:
            with self.subTest(payload=payload):
                quiz = normalize_quiz(payload, "matching", "seed")
                self.assertEqual(quiz["type"], "matching")
                self.assertEqual(len(quiz["left"]), 2)
                self.assertEqual(sorted(quiz["correct_mapping"]), [0, 1])

    def test_truncated_matching_pair_salvage(self):
        raw = '{"type":"matching","pairs":[{"term":"A","definition":"Definition A"},{"term":"B","definition":"Definition B"}'
        quiz = normalize_quiz(raw, "matching", "seed")
        self.assertEqual(len(quiz["left"]), 2)

    def test_five_complete_shapes(self):
        shapes = {
            "true_false": {"type": "true_false", "question": "A grounded statement.", "correct_index": "0"},
            "mcq": {"type": "mcq", "question": "Q", "options": ["A", "B", "C", "D"], "correct_index": "2"},
            "matching": {"type": "matching", "p": [["A", "Definition A"], ["B", "Definition B"]]},
            "fill_blank": {"type": "fill_blank", "question": "_____ is used.", "answer": "GPS"},
            "short_answer": {"type": "short_answer", "question": "Explain GPS.", "reference": "A grounded answer."},
        }
        for quiz_type, payload in shapes.items():
            with self.subTest(quiz_type=quiz_type):
                self.assertEqual(normalize_quiz(payload, quiz_type)["type"], quiz_type)


class FakeResponse:
    def __init__(self, payload=None, lines=None, content_type="application/json"):
        self.payload = payload or {}
        self.lines = lines or []
        self.headers = {"content-type": content_type}

    def raise_for_status(self):
        return None

    def json(self):
        return self.payload

    def iter_lines(self):
        return iter(self.lines)


class RoutingSession:
    def __init__(self):
        self.bot_ids = []

    def mount(self, *_args):
        pass

    def post(self, _url, **kwargs):
        bot_id = kwargs["json"]["bot_id"]
        self.bot_ids.append(bot_id)
        return FakeResponse({"content": json.dumps({"agent_bot": bot_id})})


class PersistedSession(RoutingSession):
    def post(self, _url, **kwargs):
        self.bot_ids.append(kwargs["json"]["bot_id"])
        lines = [
            'event: conversation.message.completed',
            'data: {"event":"conversation.message.completed","data":{"id":"chat-1","conversation_id":"conv-1","role":"assistant","type":"answer","content":"{\\"answer\\":\\"truncated"}}',
            'data: [DONE]',
        ]
        return FakeResponse(lines=lines, content_type="text/event-stream")

    def get(self, url, **_kwargs):
        self.last_get = url
        return FakeResponse({"data": [{"role": "assistant", "type": "answer", "content": '{"answer":"complete"}'}]})


class CompletedStreamSession(RoutingSession):
    def post(self, _url, **kwargs):
        self.bot_ids.append(kwargs["json"]["bot_id"])
        lines = [
            'event: conversation.message.completed',
            'data: {"event":"conversation.message.completed","data":{"id":"message-1","conversation_id":"conv-1","role":"assistant","type":"answer","content":"{\\"answer\\":\\"complete event\\"}"}}',
            'data: [DONE]',
        ]
        return FakeResponse(lines=lines, content_type="text/event-stream")


class NetworkFailureSession(RoutingSession):
    def __init__(self):
        super().__init__()
        self.calls = 0

    def post(self, *_args, **_kwargs):
        self.calls += 1
        raise OSError("simulated network failure")


class CozeRoutingTests(unittest.TestCase):
    def configs(self, stream=False):
        return {
            AgentKind.QUIZ: AgentConfig(AgentKind.QUIZ, "https://api.test", "token", "quiz-bot", 10, stream),
            AgentKind.GRADING: AgentConfig(AgentKind.GRADING, "https://api.test", "token", "grading-bot", 10, stream),
            AgentKind.NOTES: AgentConfig(AgentKind.NOTES, "https://api.test", "token", "notes-bot", 10, stream),
        }

    def test_three_tasks_route_to_three_bot_ids(self):
        session = RoutingSession()
        router = CozeAgentRouter(self.configs(), session)
        for kind in AgentKind:
            router.request(kind, "prompt", "student")
        self.assertEqual(session.bot_ids, ["quiz-bot", "grading-bot", "notes-bot"])

    def test_stream_truncation_reads_persisted_message(self):
        session = PersistedSession()
        router = CozeAgentRouter(self.configs(True), session)
        response = router.request(AgentKind.NOTES, "prompt", "student")
        self.assertTrue(response.recovered_from_persisted_message)
        self.assertEqual(parse_json_like(response.content)["answer"], "complete")

    def test_completed_stream_event_is_used_without_recovery(self):
        session = CompletedStreamSession()
        response = CozeAgentRouter(self.configs(True), session).request(AgentKind.NOTES, "prompt", "student")
        self.assertFalse(response.recovered_from_persisted_message)
        self.assertEqual(parse_json_like(response.content)["answer"], "complete event")

    def test_unconfigured_and_network_failure_are_bounded_errors(self):
        unconfigured = {kind: AgentConfig(kind, "https://api.test", "", "") for kind in AgentKind}
        with self.assertRaises(CozeTransportError):
            CozeAgentRouter(unconfigured, RoutingSession()).request(AgentKind.QUIZ, "prompt", "student")
        with self.assertRaises(CozeTransportError):
            CozeAgentRouter(self.configs(), NetworkFailureSession()).request(AgentKind.QUIZ, "prompt", "student")

    def test_circuit_breaker_stops_repeated_failed_network_calls(self):
        failure_session = NetworkFailureSession()
        router = CozeAgentRouter(
            self.configs(), failure_session, failure_threshold=2, cooldown_seconds=60,
        )
        for _ in range(2):
            with self.assertRaises(CozeTransportError):
                router.request(AgentKind.QUIZ, "prompt", "student")
        self.assertEqual(failure_session.calls, 2)
        self.assertTrue(router.circuit_state()["quiz"]["open"])
        with self.assertRaises(CozeTransportError):
            router.request(AgentKind.QUIZ, "prompt", "student")
        self.assertEqual(failure_session.calls, 2)


class AgentSchemaTests(unittest.TestCase):
    def test_grading_agent_schema_and_citation_allowlist(self):
        raw = json.dumps({"is_correct": "true", "score": "120", "reasoning": "Supported by [Reference 1].", "correct_answer": "GPS", "citations": [1, 99], "insufficient_evidence": "false"})
        normalized = normalize_grading(raw, "GPS", 1)
        self.assertEqual(normalized["score"], 100)
        self.assertEqual(normalized["citations"], [1])
        router = StaticRouter(AgentKind.GRADING, json.dumps({"is_correct": True, "score": 90, "reasoning": "Supported by [Reference 1].", "correct_answer": "GPS is satellite navigation.", "citations": [1], "insufficient_evidence": False}))
        result = grade_short_answer({"question": "What is GPS?", "reference": "GPS is satellite navigation."}, "Satellite navigation.", [{"source_file": "notes.pdf", "section": "GPS", "page": 62, "excerpt": "GPS is satellite navigation."}], router, "student")
        self.assertEqual(result["grading_source"], "coze_grading_agent")

    def test_notes_agent_rejects_wrong_schema_and_invalid_reference_ids(self):
        refs = [{"reference_id": 1, "source_file": "notes.pdf", "chapter": "Chapter 3", "section": "GPS", "page": 62, "excerpt": "GPS is satellite navigation.", "chunk_id": "c1", "knowledge_key": "k1"}]
        normalized = normalize_notes('{"answer":"GPS is satellite navigation.","used_reference_ids":[1,9],"insufficient_information":false}', refs)
        self.assertEqual(normalized["used_reference_ids"], [1])
        chunks = [{"source_file": "notes.pdf", "chapter": "Chapter 3", "section": "GPS", "page": 62, "text": "GPS is a satellite navigation system that provides position information worldwide.", "chunk_id": "c1", "knowledge_key": "k1"}]
        result = answer_notes("What is GPS?", chunks, StaticRouter(AgentKind.NOTES, '{"type":"mcq","question":"wrong task"}'), "student")
        self.assertEqual(result["answering_source"], "local_fallback")
        self.assertNotIn('"type"', result["answer"])


class RealPdfKnowledgeTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.temp = tempfile.TemporaryDirectory()
        cls.repo = KnowledgeRepository(Path(cls.temp.name) / "knowledge")
        cls.document = cls.repo.upload(REAL_PDF.read_bytes(), REAL_PDF.name, "automated-test")
        cls.chapter = cls.repo.chapters()[0]["id"]

    @classmethod
    def tearDownClass(cls):
        cls.temp.cleanup()

    def test_real_pdf_pages_sections_facts_and_matching_terms(self):
        self.assertEqual(self.document["page_count"], 90)
        self.assertEqual(len(self.repo.pages(self.document["id"])), 90)
        page_61 = next(item for item in self.repo.pages(self.document["id"]) if item["page"] == 61)
        self.assertEqual(page_61["section"], "Global Positioning System (GPS)")
        health = self.repo.health()
        self.assertGreater(health["chunks"], 70)
        self.assertGreater(health["facts"], 200)
        pairs = self.repo.matching_candidates("GPS components", self.chapter, 30)
        component_terms = {item["term"] for item in pairs if item["category"] == "GPS Components"}
        self.assertTrue({"GPS Satellites", "Control segment", "User (GPS Receiver)"} <= component_terms)

    def test_topic_mode_uses_outline_boundaries_and_rejects_unrelated_topics(self):
        mobile = self.repo.search("mobile communication", 20, self.chapter, {}, topic_mode=True)
        self.assertEqual(mobile["topic_scope"]["start"], 31)
        self.assertEqual(mobile["topic_scope"]["end"], 60)
        self.assertTrue(mobile["results"])
        self.assertTrue(all(31 <= int(item["page"]) <= 60 for item in mobile["results"]))
        self.assertIn(42, {int(item["page"]) for item in mobile["results"]})
        self.assertNotIn(19, {int(item["page"]) for item in mobile["results"]})

        lbs = self.repo.search("LBS", 20, self.chapter, {}, topic_mode=True)
        self.assertEqual((lbs["topic_scope"]["start"], lbs["topic_scope"]["end"]), (17, 30))
        self.assertTrue(all(17 <= int(item["page"]) <= 30 for item in lbs["results"]))

        unrelated = self.repo.search("linear algebra", 20, self.chapter, {}, topic_mode=True)
        self.assertEqual(unrelated["results"], [])
        self.assertEqual(unrelated["topic_scope"], {})

    def test_mobile_communication_short_answer_uses_network_structure_rubric(self):
        history = {"knowledge": [], "chunks": [], "pages": [], "facts": [], "fingerprints": []}
        for quiz_type in ("true_false", "mcq", "matching", "fill_blank"):
            search = self.repo.search("mobile communication", 12, self.chapter, history, topic_mode=True)
            terms = self.repo.matching_candidates("mobile communication", self.chapter, 24)
            generated = generate_quiz(
                quiz_type, "mobile communication", search, terms, history, DisabledRouter(), "student",
                generation_seed="mobile-communication-alignment",
            )
            history["knowledge"].append(generated.focus.chunk["knowledge_key"])
            history["chunks"].append(generated.focus.chunk["chunk_id"])
            history["pages"].append(generated.focus.chunk["page"])
            history["facts"].append(generated.focus.fact_key)
            history["fingerprints"].append(generated.question_fingerprint)

        short_search_history = {**history, "knowledge": [], "chunks": [], "pages": []}
        search = self.repo.search(
            "mobile communication", 20, self.chapter, short_search_history, topic_mode=True
        )
        terms = self.repo.matching_candidates("mobile communication", self.chapter, 24)
        generated = generate_quiz(
            "short_answer", "mobile communication", search, terms, history, DisabledRouter(), "student",
            generation_seed="mobile-communication-alignment",
        )
        self.assertEqual(int(generated.focus.chunk["page"]), 42)
        self.assertIn("Network structure", generated.quiz["question"])
        self.assertNotIn("points about mobile communication", generated.quiz["question"].casefold())
        expected_labels = {
            "Base Station Subsystem",
            "Network and Switching Subsystem",
            "GPRS Core Network",
            "Operations support system (OSS)",
        }
        actual_labels = {point.split(":", 1)[0] for point in generated.quiz["key_points"]}
        self.assertEqual(actual_labels, expected_labels)
        self.assertEqual({int(item["page"]) for item in generated.citations}, {42})
        validate_grounding(generated.quiz, search["results"], terms, generated.focus)

        vague_agent = StaticRouter(
            AgentKind.QUIZ,
            json.dumps({
                "type": "short_answer",
                "question": "According to the course notes, state and explain the 4 points about mobile communication?",
                "key_points": ["LBS", "Pull", "Push", "Tracking"],
                "reference": "LBS, Pull, Push and Tracking.",
                "explanation": "LBS, Pull, Push and Tracking.",
            }),
        )
        with patch.dict(os.environ, {"COZE_QUIZ_TYPES": "all"}):
            repaired = generate_quiz(
                "short_answer", "mobile communication", search, terms, history, vague_agent, "student",
                generation_seed="mobile-communication-alignment",
            )
        self.assertEqual(repaired.generation_source, "coze_quiz_agent")
        self.assertIn("Network structure", repaired.quiz["question"])
        self.assertNotIn("points about mobile communication", repaired.quiz["question"].casefold())
        self.assertEqual(
            {point.split(":", 1)[0] for point in repaired.quiz["key_points"]},
            expected_labels,
        )

    def test_short_answer_validator_rejects_cross_topic_question(self):
        search = self.repo.search("LBS", 20, self.chapter, {}, topic_mode=True)
        terms = self.repo.matching_candidates("LBS", self.chapter, 24)
        generated = generate_quiz(
            "short_answer", "LBS", search, terms, {}, DisabledRouter(), "student",
            generation_seed="short-answer-cross-topic",
        )
        invalid = dict(generated.quiz)
        invalid["question"] = "According to the course notes, state and explain the 4 points about mobile communication?"
        with self.assertRaisesRegex(QuizValidationError, "question does not match"):
            validate_grounding(invalid, search["results"], terms, generated.focus)

    def test_five_types_rotate_pages_keys_and_facts(self):
        history = {"knowledge": [], "chunks": [], "pages": [], "facts": [], "fingerprints": []}
        rows = []
        for quiz_type in ("true_false", "mcq", "matching", "fill_blank", "short_answer"):
            search = self.repo.search("GPS", 12, self.chapter, history)
            terms = self.repo.matching_candidates("GPS", self.chapter, 24)
            generated = generate_quiz(quiz_type, "GPS", search, terms, history, DisabledRouter(), "student")
            row = (quiz_type, generated.focus.chunk["page"], generated.focus.chunk["knowledge_key"], generated.focus.fact_key, generated.focus.text)
            rows.append(row)
            history["knowledge"].append(row[2]); history["chunks"].append(generated.focus.chunk["chunk_id"])
            history["pages"].append(row[1]); history["facts"].append(row[3]); history["fingerprints"].append(generated.question_fingerprint)
            self.assertGreater(lexical_score("GPS", f"{generated.focus.chunk['section']} {generated.focus.text}"), 0)
        self.generation_rows = rows
        self.assertGreater(len({row[1] for row in rows}), 1)
        self.assertGreater(len({row[2] for row in rows}), 1)
        self.assertEqual(len({row[3] for row in rows}), 5)
        broad = self.repo.search("GPS", 20, self.chapter, {})["results"]
        self.assertIn(71, {item["page"] for item in broad})

    def test_fill_blank_rejects_grammar_word_and_uses_course_concept(self):
        search = self.repo.search("GPS", 20, self.chapter, {}, topic_mode=True)
        terms = self.repo.matching_candidates("GPS", self.chapter, 24)
        source_chunk = next(
            chunk for chunk in search["results"]
            if any("For the GPS user" in fact.get("text", "") for fact in chunk.get("facts", []))
        )
        source_fact = next(
            fact for fact in source_chunk["facts"] if "For the GPS user" in fact.get("text", "")
        )
        chunk = dict(source_chunk)
        chunk["facts"] = [source_fact]
        chunk["fact_keys"] = [source_fact["fact_key"]]
        focus = FocusFact(source_fact["text"], source_fact["fact_key"], chunk)

        grammar_item = {
            "type": "fill_blank",
            "question": "_____ the GPS user, the system is passive and requires a receiver only without the requirement to transmit.",
            "answer": "For",
            "explanation": source_fact["text"],
        }
        with self.assertRaisesRegex(QuizValidationError, "grammar/function word"):
            validate_grounding(grammar_item, [chunk], terms, focus)

        local, _citations = local_quiz("fill_blank", focus, [chunk], terms, "GPS")
        self.assertEqual(local["answer"].casefold(), "receiver")
        self.assertIn("requires a _____ only", local["question"])
        validate_grounding(local, [chunk], terms, focus)

        bad_agent = StaticRouter(AgentKind.QUIZ, json.dumps(grammar_item))
        with patch.dict(os.environ, {"COZE_QUIZ_TYPES": "all"}):
            repaired = generate_quiz(
                "fill_blank", "GPS", {"results": [chunk]}, terms, {}, bad_agent, "student",
                generation_seed="fill-blank-function-word",
            )
        self.assertEqual(repaired.generation_source, "local_fallback")
        self.assertEqual(repaired.quiz["answer"].casefold(), "receiver")
        self.assertNotEqual(repaired.quiz["question"].lstrip()[:5], "_____")

    def test_fill_blank_fallback_survives_multiple_mcq_retries(self):
        history = {"knowledge": [], "chunks": [], "pages": [], "facts": [], "fingerprints": []}
        # This sequence reproduces the demo failure from v8.0.7: several MCQ
        # retries consume the easiest GPS facts before Fill Blank is reached.
        for quiz_type in ("true_false", "mcq", "mcq", "mcq", "matching"):
            search = self.repo.search("GPS", 12, self.chapter, history)
            terms = self.repo.matching_candidates("GPS", self.chapter, 24)
            generated = generate_quiz(
                quiz_type, "GPS", search, terms, history, DisabledRouter(), "student",
                generation_seed="fill-blank-regression-session",
            )
            history["knowledge"].append(generated.focus.chunk["knowledge_key"])
            history["chunks"].append(generated.focus.chunk["chunk_id"])
            history["pages"].append(generated.focus.chunk["page"])
            history["facts"].append(generated.focus.fact_key)
            history["fingerprints"].append(generated.question_fingerprint)

        search = self.repo.search("GPS", 12, self.chapter, history)
        terms = self.repo.matching_candidates("GPS", self.chapter, 24)
        generated = generate_quiz(
            "fill_blank", "GPS", search, terms, history, DisabledRouter(), "student",
            generation_seed="fill-blank-regression-session",
        )
        self.assertIn("_____", generated.quiz["question"])
        self.assertNotIn("..", generated.quiz["question"])
        self.assertTrue(generated.quiz["answer"])
        self.assertIn(generated.quiz["answer"].casefold(), generated.focus.text.casefold())
        validate_grounding(generated.quiz, search["results"], terms, generated.focus)

    def test_true_false_uses_both_answer_indices_across_session_seeds(self):
        search = self.repo.search("GPS", 12, self.chapter, {})
        terms = self.repo.matching_candidates("GPS", self.chapter, 24)
        generated = [
            generate_quiz(
                "true_false", "GPS", search, terms, {}, DisabledRouter(), "student",
                generation_seed=f"session-{index}",
            )
            for index in range(12)
        ]
        self.assertEqual({item.quiz["correct_index"] for item in generated}, {0, 1})
        for item in generated:
            validate_grounding(item.quiz, search["results"], terms, item.focus)
            self.assertEqual(
                item.quiz["explanation"].rstrip(". ").casefold(),
                item.focus.text.rstrip(". ").casefold(),
            )
            if item.quiz["correct_index"] == 1:
                self.assertNotEqual(item.quiz["question"], item.quiz["explanation"])

    def test_gps_false_item_uses_a_complete_bounded_contradiction(self):
        search = self.repo.search("gps", 12, self.chapter, {}, topic_mode=True)
        terms = self.repo.matching_candidates("gps", self.chapter, 24)
        false_seed = next(
            seed for seed in (f"gps-false-{index}" for index in range(100))
            if generate_quiz(
                "true_false", "gps", search, terms, {}, DisabledRouter(), "student",
                generation_seed=seed,
            ).quiz["correct_index"] == 1
        )
        generated = generate_quiz(
            "true_false", "gps", search, terms, {}, DisabledRouter(), "student",
            generation_seed=false_seed,
        )
        self.assertEqual(generated.quiz["correct_index"], 1)
        self.assertFalse(generated.quiz["question"].casefold().startswith("the course notes do not state"))
        self.assertFalse(generated.focus.text.casefold().startswith("based on "))
        self.assertGreaterEqual(len(generated.focus.text.split()), 6)
        validate_grounding(generated.quiz, search["results"], terms, generated.focus)

    def test_true_false_agent_cannot_ignore_required_balance_target(self):
        search = self.repo.search("GPS", 12, self.chapter, {})
        terms = self.repo.matching_candidates("GPS", self.chapter, 24)
        false_seed = next(
            seed for seed in (f"forced-false-{index}" for index in range(100))
            if generate_quiz(
                "true_false", "GPS", search, terms, {}, DisabledRouter(), "student",
                generation_seed=seed,
            ).quiz["correct_index"] == 1
        )
        always_true_agent = StaticRouter(
            AgentKind.QUIZ,
            json.dumps({
                "type": "true_false",
                "question": "GPS is a satellite based navigation system offering precision navigation capability.",
                "options": ["True", "False"],
                "correct_index": 0,
                "explanation": "GPS is a satellite based navigation system offering precision navigation capability.",
            }),
        )
        generated = generate_quiz(
            "true_false", "GPS", search, terms, {}, always_true_agent, "student",
            generation_seed=false_seed,
        )
        self.assertEqual(generated.generation_source, "local_fallback")
        self.assertEqual(generated.quiz["correct_index"], 1)

    def test_invalid_coze_mcq_is_rejected_and_replaced_by_grounded_local_question(self):
        bad_agent = StaticRouter(
            AgentKind.QUIZ,
            json.dumps({
                "type": "mcq",
                "question": "Which course-note term correctly completes this statement? _____ as a definer/launcher for area alarm.",
                "options": ["to restrict results of search that", "Use", "Class discussions: What", "Class discussions"],
                "correct_index": 1,
                "explanation": "Use as a definer/launcher for area alarm.",
            }),
        )
        search = self.repo.search("LBS", 20, self.chapter, {})
        generated = generate_quiz(
            "mcq", "LBS", search, self.repo.matching_candidates("LBS", self.chapter, 24),
            {}, bad_agent, "student",
        )
        self.assertEqual(generated.generation_source, "local_fallback")
        self.assertNotIn("correctly completes this statement", generated.quiz["question"].casefold())
        self.assertTrue(generated.quiz["question"].endswith("?"))
        self.assertEqual(len(generated.quiz["options"]), 4)
        self.assertNotIn("Use", generated.quiz["options"])
        self.assertNotIn("Class discussions: What", generated.quiz["options"])

    def test_lbs_mcq_is_a_natural_grounded_question_not_a_fragment_completion(self):
        search = self.repo.search("LBS", 20, self.chapter, {})
        chunks = search["results"]
        page_22 = next(item for item in chunks if item["page"] == 22)
        fact = next(item for item in page_22["facts"] if "definer/launcher" in item["text"])
        focus = FocusFact(fact["text"], fact["fact_key"], page_22)
        terms = self.repo.matching_candidates("LBS", self.chapter, 24)
        quiz, _citations = local_quiz("mcq", focus, chunks, terms, "LBS")
        validate_grounding(quiz, chunks, terms, focus)
        self.assertTrue(quiz["question"].endswith("?"))
        self.assertNotIn("correctly completes this statement", quiz["question"].casefold())
        self.assertNotIn("_____", quiz["question"])
        for bad in ("Use", "Class discussions: What", "to restrict results of search that"):
            self.assertNotIn(bad, quiz["options"])
        self.assertIn("area alarm", quiz["question"].casefold())
        self.assertIn("launch a pre-defined activity", quiz["options"][quiz["correct_index"]])

    def test_notes_and_objective_grading(self):
        refs = [{"reference_id": 1, "excerpt": "GPS is a satellite based navigation system offering precision navigation capability."}]
        answer = local_notes_answer("What is GPS?", refs)
        self.assertFalse(answer["insufficient_information"])
        missing = local_notes_answer("What is the cafeteria menu?", refs)
        self.assertEqual(missing["answer"], NOT_FOUND)
        result = grade_objective({"type": "mcq", "options": ["A", "B", "C", "D"], "correct_index": 1}, "B")
        self.assertTrue(result["is_correct"])
        semantic = local_semantic_grade("satellite navigation system", "GPS is a satellite based navigation system", 1)
        self.assertEqual(semantic["grading_source"], "local_fallback")

    def test_agent_failure_fast_falls_back_to_grounded_local_quiz(self):
        search = self.repo.search("GPS", 12, self.chapter, {})
        terms = self.repo.matching_candidates("GPS", self.chapter, 24)
        generated = generate_quiz("true_false", "GPS", search, terms, {}, FailingRouter(), "student")
        self.assertEqual(generated.generation_source, "local_fallback")
        self.assertEqual(
            generated.focus.text.rstrip(". ").casefold(),
            generated.quiz["explanation"].rstrip(". ").casefold(),
        )
        self.assertIn(generated.quiz["correct_index"], {0, 1})

    def test_knowledge_persists_after_repository_instance_closes(self):
        reopened = KnowledgeRepository(Path(self.temp.name) / "knowledge")
        self.assertEqual(len(reopened.documents()), 1)
        results = reopened.search("GPS", 8, self.chapter, {})["results"]
        self.assertTrue(results)
        self.assertTrue(any("GPS" in f"{item['section']} {item['text']}" for item in results))


    def test_exhausted_history_can_reuse_grounded_evidence_for_every_type(self):
        search = self.repo.search("GPS", 20, self.chapter, {}, topic_mode=True)
        terms = self.repo.matching_candidates("GPS", self.chapter, 30)
        exhausted_history = {
            "knowledge": [item.get("knowledge_key") for item in search["results"]],
            "chunks": [item.get("chunk_id") for item in search["results"]],
            "pages": [item.get("page") for item in search["results"]],
            "facts": [
                fact.get("fact_key")
                for item in search["results"]
                for fact in item.get("facts", [])
            ],
            "fingerprints": [f"Q-used-{index}" for index in range(100)],
        }
        for quiz_type in ("true_false", "mcq", "matching", "fill_blank", "short_answer"):
            with self.subTest(quiz_type=quiz_type):
                generated = generate_quiz(
                    quiz_type, "GPS", search, terms, exhausted_history, DisabledRouter(), "student",
                    generation_seed=f"exhausted-{quiz_type}", allow_agent=False, local_source="local_recovery",
                )
                self.assertEqual(generated.generation_source, "local_recovery")
                validate_grounding(generated.quiz, search["results"], terms, generated.focus)


class TeacherHybridSearchTests(unittest.TestCase):
    def test_hybrid_search_recovers_acronyms_paraphrases_and_minor_typos(self):
        records = [
            SimpleNamespace(
                id="1", student_name="Alice Wong", topic="GPS", chapter="Mobile technologies",
                question_type="short_answer",
                question="Explain how the Global Positioning System supports satellite navigation.",
                answer="It determines position from satellite signals.", correct_answer="Satellite navigation",
                ai_feedback="", teacher_feedback="", explanation="", update_time="2026-08-01", created_at="2026-08-01",
            ),
            SimpleNamespace(
                id="2", student_name="Bob Chan", topic="LBS", chapter="Location services",
                question_type="mcq", question="What is a location-based service?",
                answer="A service using geographical position.", correct_answer="Location-based service",
                ai_feedback="", teacher_feedback="", explanation="", update_time="2026-08-01", created_at="2026-08-01",
            ),
        ]
        cases = {
            "GPS": "1",
            "global positioning system": "1",
            "satellite navigation": "1",
            "locatoin based service": "2",
            "Alice": "1",
        }
        for query, expected_id in cases.items():
            with self.subTest(query=query):
                ranked = hybrid_search_records(records, query, limit=5)
                self.assertTrue(ranked)
                self.assertEqual(ranked[0][0].id, expected_id)
                self.assertIn("search_backend", ranked[0][1])
        self.assertEqual(hybrid_search_records(records, "unrelated cafeteria menu", limit=5), [])


class StartupRegressionTests(unittest.TestCase):
    def test_blank_boolean_environment_values_use_declared_default(self):
        previous = os.environ.get("ALLOW_TEACHER_REGISTRATION")
        try:
            os.environ["ALLOW_TEACHER_REGISTRATION"] = ""
            self.assertTrue(env_bool("ALLOW_TEACHER_REGISTRATION", True))
        finally:
            if previous is None:
                os.environ.pop("ALLOW_TEACHER_REGISTRATION", None)
            else:
                os.environ["ALLOW_TEACHER_REGISTRATION"] = previous

    def test_blank_secret_values_still_enable_flask_sessions(self):
        previous = {
            name: os.environ.get(name)
            for name in ("APP_SECRET_KEY", "TEACHER_SECRET_KEY", "STUDENT_SECRET_KEY")
        }
        try:
            os.environ["APP_SECRET_KEY"] = ""
            os.environ["TEACHER_SECRET_KEY"] = ""
            os.environ["STUDENT_SECRET_KEY"] = ""
            self.assertTrue(app_secret())
            self.assertTrue(service_secret("TEACHER_SECRET_KEY"))
            self.assertTrue(service_secret("STUDENT_SECRET_KEY"))
        finally:
            for name, value in previous.items():
                if value is None:
                    os.environ.pop(name, None)
                else:
                    os.environ[name] = value

    def test_no_env_file_uses_local_mode_despite_stale_coze_process_values(self):
        self.assertFalse((ROOT / ".env").is_file())
        process_env = os.environ.copy()
        process_env.update({
            "PYTHONPATH": str(ROOT),
            "COZE_TOKEN": "stale-windows-token",
            "COZE_QUIZ_BOT_ID": "stale-windows-bot",
        })
        process_env.pop("COZE_ENABLED", None)
        result = subprocess.run(
            [
                os.sys.executable,
                "-c",
                "from common.config import coze_runtime_enabled, env_source; "
                "print(coze_runtime_enabled()); print(env_source('COZE_TOKEN'))",
            ],
            cwd=ROOT,
            env=process_env,
            capture_output=True,
            text=True,
            check=False,
        )
        self.assertEqual(result.returncode, 0, result.stderr)
        self.assertEqual(result.stdout.strip().splitlines(), ["False", "process_environment"])

    def test_student_javascript_parses_and_forms_are_post(self):
        node = shutil.which("node")
        if node:
            result = subprocess.run(
                [node, "--check", str(ROOT / "student_app" / "static" / "app.js")],
                capture_output=True,
                text=True,
                check=False,
            )
            self.assertEqual(result.returncode, 0, result.stderr)

        for name in ("login.html", "register.html", "exercise_setup.html"):
            html = (ROOT / "student_app" / "templates" / name).read_text(encoding="utf-8")
            self.assertIn('method="post"', html)

        self.assertTrue((ROOT / "student_app" / "static" / "images" / "ece-logo.jpg").is_file())


    def test_legacy_sqlite_schema_upgrade_preserves_accounts_and_history(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "legacy.db"
            engine = create_engine(f"sqlite:///{path}")
            with engine.begin() as connection:
                connection.execute(text("""
                    CREATE TABLE users (
                      id INTEGER PRIMARY KEY AUTOINCREMENT,
                      username VARCHAR(80) UNIQUE NOT NULL,
                      password_hash VARCHAR(256) NOT NULL,
                      name VARCHAR(80) NOT NULL DEFAULT '',
                      role VARCHAR(20) NOT NULL DEFAULT 'student'
                    )
                """))
                connection.execute(text("""
                    CREATE TABLE student_qa_records (
                      id VARCHAR(64) PRIMARY KEY, student_name VARCHAR(128) NOT NULL,
                      chapter VARCHAR(128) DEFAULT '', question_type VARCHAR(32) NOT NULL,
                      topic VARCHAR(256) DEFAULT '', question TEXT NOT NULL, answer TEXT NOT NULL,
                      options_or_pairs TEXT, correct_answer TEXT, is_correct BOOLEAN DEFAULT 0,
                      score INTEGER DEFAULT 0, ai_feedback TEXT, teacher_feedback TEXT, explanation TEXT,
                      source VARCHAR(16) DEFAULT 'coze', teacher_score INTEGER, teacher_verified BOOLEAN DEFAULT 0,
                      verified_by VARCHAR(128) DEFAULT '', verified_at VARCHAR(32) DEFAULT '',
                      time VARCHAR(32) DEFAULT '', update_time VARCHAR(32) DEFAULT '',
                      created_at DATETIME DEFAULT CURRENT_TIMESTAMP
                    )
                """))
                connection.execute(text("INSERT INTO users(username,password_hash,name,role) VALUES ('test1','legacy-hash-must-not-change','Test One','student')"))
                connection.execute(text("INSERT INTO student_qa_records(id,student_name,question_type,question,answer) VALUES ('old-record','test1','mcq','Old question','Old answer')"))

            class FakeDB:
                pass

            fake = FakeDB()
            fake.engine = engine
            fake.session = sessionmaker(bind=engine)()
            changes = ensure_legacy_schema(fake)
            self.assertIn("users.created_at", changes)
            with engine.connect() as connection:
                account = connection.execute(text("SELECT username,password_hash,role,created_at FROM users WHERE username='test1'")) .mappings().one()
                record = connection.execute(text("SELECT id,grading_source,citations_json FROM student_qa_records WHERE id='old-record'")) .mappings().one()
            self.assertEqual(account["password_hash"], "legacy-hash-must-not-change")
            self.assertEqual(account["role"], "student")
            self.assertIsNotNone(account["created_at"])
            self.assertEqual(record["id"], "old-record")
            self.assertEqual(record["grading_source"], "local")
            self.assertEqual(record["citations_json"], "[]")

    def test_knowledge_service_is_api_only_with_health_status(self):
        source = (ROOT / "knowledge_service" / "app.py").read_text(encoding="utf-8")
        self.assertIn('@app.get("/")', source)
        self.assertIn("API-only Knowledge Service", source)
        self.assertNotIn("render_template", source)

    def test_original_teacher_ui_and_registration_are_preserved(self):
        expected = {
            "login.html": "0488f9be8dea7d0f90ad82b57a80cca46d3ec172caa7a2b9435ed7649ac8ac25",
            "register.html": "9ea002e7e12c117ed1c8bbb2e0ada564413f4150356ed64fab46a542422eeeea",
            "dashboard.html": "3d7c39737074ab966687f621382dd79ab94d9d4288aafe831242e47f4bb8e8a3",
        }
        for name, digest in expected.items():
            payload = (ROOT / "teacher_app" / "templates" / name).read_bytes()
            self.assertEqual(hashlib.sha256(payload).hexdigest(), digest, name)
        login = (ROOT / "teacher_app" / "templates" / "login.html").read_text(encoding="utf-8")
        self.assertIn('href="/register"', login)
        teacher_source = (ROOT / "teacher_app" / "app.py").read_text(encoding="utf-8")
        self.assertIn('@app.route("/register", methods=["GET", "POST"])', teacher_source)
        self.assertIn('"SESSION_COOKIE_NAME": "kbot_teacher_session"', teacher_source)
        self.assertIn('"LEGACY_TEACHER_COMPAT"', teacher_source)
        self.assertIn('password = app.config["DEMO_TEACHER_PASSWORD"] or "123456"', teacher_source)

    def test_original_teacher_dashboard_endpoints_remain_available(self):
        dashboard = (ROOT / "teacher_app" / "templates" / "dashboard.html").read_text(encoding="utf-8")
        teacher_source = (ROOT / "teacher_app" / "app.py").read_text(encoding="utf-8")
        expected_routes = (
            "/dashboard", "/upload", "/knowledge/upload", "/sync_from_students",
            "/knowledge/list", "/knowledge/delete", "/update", "/batch_update",
            "/ai_reevaluate", "/batch_delete", "/search",
            "/search_by_student_name", "/delete",
        )
        for route in expected_routes:
            self.assertIn(route, dashboard, route)
            self.assertIn(f'"{route}"', teacher_source, route)
        self.assertIn('/get_detail/<record_id>', teacher_source)
        self.assertIn('/student_profile/<path:name>', teacher_source)

    def test_student_moodle_ui_and_locked_sequential_progress(self):
        template = (ROOT / "student_app" / "templates" / "exercise_quiz.html").read_text(encoding="utf-8")
        self.assertNotIn('<a class="flow-step', template)
        self.assertIn('<div class="flow-step', template)
        self.assertIn("The same topic is used for all five question types.", template)
        self.assertIn("current_flow_index", template)
        base = (ROOT / "student_app" / "templates" / "base.html").read_text(encoding="utf-8")
        self.assertIn("hku-crest.jpg", base)
        self.assertIn("ece-logo.jpg", base)
        student_source = (ROOT / "student_app" / "app.py").read_text(encoding="utf-8")
        self.assertIn('"SESSION_COOKIE_NAME": "kbot_student_session"', student_source)

    def test_answers_and_references_are_hidden_until_submission_and_history_returns_to_quiz(self):
        exercise_source = (ROOT / "student_app" / "blueprints" / "exercise.py").read_text(encoding="utf-8")
        self.assertIn('"quiz": _public_quiz(existing["quiz"])', exercise_source)
        public_block = exercise_source[exercise_source.index("def _public_quiz"):exercise_source.index("def _loads")]
        for hidden in ("correct_index", "correct_mapping", "answer", "reference", "explanation", "citations"):
            self.assertNotIn(f'public["{hidden}"]', public_block)
        javascript = (ROOT / "student_app" / "static" / "app.js").read_text(encoding="utf-8")
        generation_block = javascript[javascript.index("async function generateQuestion"):javascript.index('form.addEventListener("submit"')]
        self.assertNotIn("data.citations", generation_block)
        self.assertIn('renderCitations("#citations", result.citations)', javascript)
        self.assertIn("Reference answer:", javascript)
        history = (ROOT / "student_app" / "templates" / "history.html").read_text(encoding="utf-8")
        base = (ROOT / "student_app" / "templates" / "base.html").read_text(encoding="utf-8")
        self.assertIn("back_to_quiz_url", history)
        self.assertIn("Back to Quiz", (ROOT / "student_app" / "blueprints" / "history.py").read_text(encoding="utf-8"))
        self.assertIn("back-to-quiz-link", base)

    def test_generation_and_grading_have_anti_stall_recovery_contracts(self):
        exercise_source = (ROOT / "student_app" / "blueprints" / "exercise.py").read_text(encoding="utf-8")
        javascript = (ROOT / "student_app" / "static" / "app.js").read_text(encoding="utf-8")
        self.assertIn('@bp.post("/api/exercise/skip")', exercise_source)
        self.assertIn('"pending": True', exercise_source)
        self.assertIn('allow_agent=not recovery', exercise_source)
        self.assertIn('local_source="local_recovery"', exercise_source)
        self.assertIn('local_emergency_fallback', exercise_source)
        self.assertIn('showSkipAction', javascript)
        self.assertIn('data.pending', javascript)
        self.assertIn('/api/exercise/skip', javascript)

    def test_teacher_search_route_uses_hybrid_ranker_not_keyword_only_sql(self):
        source = (ROOT / "teacher_app" / "app.py").read_text(encoding="utf-8")
        dashboard = (ROOT / "teacher_app" / "templates" / "dashboard.html").read_text(encoding="utf-8")
        search_block = source[source.index('def search_records'):source.index('@app.post("/search_by_student_name")')]
        self.assertIn('hybrid_search_records(', search_block)
        self.assertNotIn('.ilike(', search_block)
        self.assertIn('Hybrid semantic + keyword search', dashboard)

    def test_representative_teacher_and_student_templates_render(self):
        def fake_url_for(endpoint, **values):
            if endpoint == "static":
                return "/static/" + values.get("filename", "")
            routes = {
                "portal.dashboard": "/dashboard",
                "exercise.setup_page": "/exercise/setup",
                "notes.notes_page": "/notes",
                "history.history_page": "/history",
                "auth.login_action": "/login",
                "auth.register_page": "/register",
                "auth.register_action": "/register",
                "auth.login_page": "/login",
                "exercise.setup_action": "/api/exercise/setup",
            }
            return routes.get(endpoint, "/" + endpoint.replace(".", "/"))

        student_env = Environment(loader=FileSystemLoader(ROOT / "student_app" / "templates"))
        student_env.globals["url_for"] = fake_url_for
        shared = {
            "session": {"student_username": "test1", "display_name": "Test One"},
            "config": {"ENABLED_STUDENT_MODULES": {"exercise", "notes", "history"}},
        }
        html = student_env.get_template("exercise_quiz.html").render(
            **shared,
            label="MCQ",
            chapter="Chapter 3",
            quiz_type="mcq",
            labels={
                "true_false": "True/False", "mcq": "MCQ", "matching": "Matching",
                "fill_blank": "Fill Blank", "short_answer": "Short Answer",
            },
            types=("true_false", "mcq", "matching", "fill_blank", "short_answer"),
            current_flow_index=1, topic="GPS", topic_locked=True, auto_generate=True,
        )
        self.assertIn("Topic for this quiz flow", html)
        self.assertNotIn('<a class="flow-step', html)
        self.assertNotIn('id="topic"', html)

        teacher_env = Environment(loader=FileSystemLoader(ROOT / "teacher_app" / "templates"))
        teacher_env.globals["url_for"] = fake_url_for
        self.assertIn("Create one", teacher_env.get_template("login.html").render())
        dashboard = teacher_env.get_template("dashboard.html").render(
            items=[], page=1, total_pages=1, total_items=0, user={"name": "Teacher"}
        )
        self.assertIn("knowledgeInput", dashboard)

    def test_sequential_flow_transition_contract(self):
        order = ("true_false", "mcq", "matching", "fill_blank", "short_answer")
        current = order[0]
        for expected in order[1:]:
            wrong = transition(current, False)
            self.assertEqual(wrong.next_type, current)
            self.assertFalse(wrong.completed)
            right = transition(current, True)
            self.assertEqual(right.next_type, expected)
            self.assertTrue(right.advanced)
            current = expected
        final = transition("short_answer", False)
        self.assertTrue(final.completed)
        self.assertIsNone(final.next_type)
        skipped = skip_transition("matching")
        self.assertEqual(skipped.next_type, "fill_blank")
        self.assertTrue(skipped.advanced)
        skipped_final = skip_transition("short_answer")
        self.assertTrue(skipped_final.completed)
        self.assertFalse(request_matches_current("short_answer", "mcq"))

    def test_docker_reuses_v62_mysql_volume_and_account_check_exists(self):
        compose = (ROOT / "docker-compose.yml").read_text(encoding="utf-8")
        self.assertIn("edu_qa_mysql_data", compose)
        self.assertIn("MYSQL_VOLUME_NAME", compose)
        self.assertTrue((ROOT / "CHECK_ACCOUNTS.bat").is_file())
        self.assertTrue((ROOT / "scripts" / "account_diagnostics.py").is_file())

class KahootExportTests(unittest.TestCase):
    def test_kahoot_xlsx_uses_official_template_and_skips_malformed_records(self):
        from openpyxl import load_workbook

        records = [
            SimpleNamespace(
                question_type="mcq",
                question="Which statement describes GPS?",
                options_or_pairs=json.dumps({"options": ["A system", "Satellite navigation", "A sensor", "A database"]}),
                correct_answer="Satellite navigation",
            ),
            SimpleNamespace(
                question_type="mcq", question="Malformed old record",
                options_or_pairs="", correct_answer="",
            ),
        ]
        template = ROOT / "teacher_app" / "kahoot_template.xlsx"
        self.assertTrue(template.is_file())
        output, count = build_kahoot_xlsx(records, template)
        self.assertEqual(count, 1)
        workbook = load_workbook(output)
        sheet = workbook.active
        self.assertEqual(sheet.cell(8, 1).value, "Question - max 120 characters")
        self.assertEqual(sheet.cell(9, 1).value, "Which statement describes GPS?")
        self.assertEqual(sheet.cell(9, 3).value, "Satellite navigation")
        self.assertEqual(str(sheet.cell(9, 7).value), "2")
        self.assertEqual(len(eligible_questions(records)), 1)

    def test_teacher_route_exports_xlsx_and_has_safe_fallback(self):
        source = (ROOT / "teacher_app" / "app.py").read_text(encoding="utf-8")
        requirements = (ROOT / "requirements.txt").read_text(encoding="utf-8").casefold()
        self.assertIn("build_kahoot_xlsx", source)
        self.assertIn('download_name="kahoot_quiz_template.xlsx"', source)
        self.assertIn("Kahoot XLSX export failed", source)
        self.assertIn("openpyxl", requirements)


class RelaxedKnowledgeRuleTests(unittest.TestCase):
    def test_mcq_allows_synthetic_distractors_when_correct_option_is_grounded(self):
        correct = "In building services, the lamp is the electrical component that produces light."
        chunk = {
            "source_file": "Intro 2026.pdf",
            "chapter": "Lighting",
            "section": "Vocabulary",
            "page": 18,
            "text": correct,
            "chunk_id": "lighting-18",
            "knowledge_key": "lighting-vocabulary",
            "facts": [{"text": correct, "fact_key": "lamp-definition"}],
        }
        focus = FocusFact(correct, "lamp-definition", chunk)
        quiz = {
            "type": "mcq",
            "question": "Which statement best describes a lamp in building services?",
            "options": [
                correct,
                "It is the enclosure that stores electrical energy for the building.",
                "It is a sensor that measures the colour temperature of daylight.",
                "It is a controller that converts illuminance directly into radiant power.",
            ],
            "correct_index": 0,
            "explanation": correct,
        }
        validate_grounding(quiz, [chunk], [], focus)

    def test_mcq_still_rejects_an_ungrounded_correct_option(self):
        note = "In building services, the lamp is the electrical component that produces light."
        chunk = {
            "text": note,
            "section": "Vocabulary",
            "facts": [{"text": note, "fact_key": "lamp-definition"}],
        }
        focus = FocusFact(note, "lamp-definition", chunk)
        quiz = {
            "type": "mcq",
            "question": "Which statement best describes a lamp in building services?",
            "options": [
                "A lamp is a device that measures illuminance.",
                "It stores electrical energy.",
                "It is a daylight sensor.",
                "It controls room temperature.",
            ],
            "correct_index": 0,
            "explanation": note,
        }
        with self.assertRaises(QuizValidationError):
            validate_grounding(quiz, [chunk], [], focus)

    def test_narrow_topic_local_mcq_can_use_broader_retrieved_facts_as_distractors(self):
        correct = "In building services, the lamp is the electrical component that produces light."
        focus_chunk = {
            "text": correct,
            "section": "Vocabulary",
            "chunk_id": "c1",
            "facts": [{"text": correct, "fact_key": "lamp-definition"}],
        }
        other_chunks = [
            {
                "text": text,
                "section": section,
                "chunk_id": f"c{index}",
                "facts": [{"text": text, "fact_key": f"f{index}"}],
            }
            for index, (section, text) in enumerate([
                ("Vision", "Rods are more sensitive than cones under low light conditions."),
                ("Spectrum", "Visible light is a small part of the electromagnetic spectrum."),
                ("Adaptation", "Dark adaptation requires more time than light adaptation."),
            ], 2)
        ]
        focus = FocusFact(correct, "lamp-definition", focus_chunk)
        quiz, _citations = local_quiz(
            "mcq", focus, [focus_chunk, *other_chunks], [], "definition of lamp"
        )
        self.assertEqual(len(quiz["options"]), 4)
        self.assertIn(correct, quiz["options"])
        self.assertIn("definition of lamp", quiz["question"].casefold())
        validate_grounding(quiz, [focus_chunk, *other_chunks], [], focus)

    def test_notes_agent_may_answer_with_general_knowledge_without_citations(self):
        raw = json.dumps({
            "answer": "A concise general-knowledge explanation.",
            "used_reference_ids": [],
            "insufficient_information": False,
        })
        normalized = normalize_notes(raw, [])
        self.assertEqual(normalized["used_reference_ids"], [])
        self.assertFalse(normalized["insufficient_information"])

        result = answer_notes(
            "Explain this concept.",
            [],
            StaticRouter(AgentKind.NOTES, raw),
            "student",
        )
        self.assertEqual(result["answering_source"], "coze_notes_agent")
        self.assertEqual(result["citations"], [])
        self.assertEqual(result["answer"], "A concise general-knowledge explanation.")


class V3KnowledgeRuleRegressionTests(unittest.TestCase):
    @staticmethod
    def _chunk(fact, topic="Bayes risk"):
        return {
            "text": fact,
            "section": topic,
            "page": 18,
            "source_file": "BayesDecisionTheory.docx",
            "chunk_id": "bayes-test",
            "knowledge_key": "bayes-knowledge",
            "facts": [{"text": fact, "fact_key": "bayes-fact"}],
        }

    def test_tf_retry_alternates_answers_for_same_session_and_topic(self):
        from common.quiz_engine import _true_false_target
        fact = "The classifier minimizes Bayes risk under the chosen loss function."
        chunk = self._chunk(fact)
        history = {"facts": [], "fingerprints": []}
        actual = []
        for _ in range(4):
            result = generate_quiz(
                "true_false", "Bayes risk", {"results": [chunk]}, [], history,
                DisabledRouter(), "student", generation_seed="same-session",
                allow_agent=False,
            )
            actual.append(result.quiz["correct_index"])
            history["facts"].append(result.focus.fact_key)
            history["fingerprints"].append(result.question_fingerprint)
            validate_grounding(result.quiz, [chunk], [], result.focus)
        self.assertEqual(actual, [
            _true_false_target("Bayes risk", "same-session", i) for i in range(4)
        ])
        self.assertEqual(actual[0], actual[2])
        self.assertNotEqual(actual[0], actual[1])

    def test_tf_uses_supported_true_if_false_cannot_be_constructed(self):
        from common.quiz_engine import _true_false_target
        fact = "Conditional risk combines losses using posterior class probabilities."
        chunk = self._chunk(fact, "conditional risk")
        false_seed = next(
            f"no-false-{i}" for i in range(100)
            if _true_false_target("conditional risk", f"no-false-{i}") == 1
        )
        result = generate_quiz(
            "true_false", "conditional risk", {"results": [chunk]}, [], {},
            DisabledRouter(), "student", generation_seed=false_seed,
            allow_agent=False,
        )
        self.assertEqual(result.quiz["correct_index"], 0)
        self.assertIn("combines losses", result.quiz["question"])
        validate_grounding(result.quiz, [chunk], [], result.focus)

    def test_short_bayes_proposition_can_be_falsified_safely(self):
        from common.quiz_engine import _false_variant, _bounded_false_contradiction
        source = "The classifier minimizes Bayes risk."
        changed = _false_variant(source, "seed")
        self.assertEqual(changed, "The classifier maximizes Bayes risk.")
        self.assertTrue(_bounded_false_contradiction(changed, source))

    def test_mcq_accepts_readable_longer_technical_option(self):
        from common.quiz_engine import _mcq_option_quality
        correct = (
            "Under equal class variance and independent features, a classifier "
            "with uniform prior probabilities can use the distance to a class "
            "mean; the relevant decision criterion is obtained from the "
            "class-conditional probability and the prior probability."
        )
        self.assertTrue(180 < len(correct) <= 320)
        self.assertTrue(_mcq_option_quality(correct))
        chunk = self._chunk(correct, "minimum distance classifier")
        focus = FocusFact(correct, "bayes-fact", chunk)
        quiz = {
            "type":"mcq",
            "question":"Which statement describes the minimum-distance classifier?",
            "options":[
                correct,
                "It assigns classes by reversing the sign of each prior probability.",
                "It requires a larger posterior probability for every other class.",
                "It chooses classes by ignoring both the class means and the observations."
            ],
            "correct_index":0,
            "explanation":correct,
        }
        validate_grounding(quiz,[chunk],[],focus)
        self.assertFalse(_mcq_option_quality("Class discussions: What"))
        self.assertFalse(_mcq_option_quality("P(x | (class)"))

    def test_notes_prompt_enforces_complete_conditions_and_safe_formulas(self):
        from common.notes_engine import notes_prompt
        prompt = notes_prompt("Under what conditions is minimum distance valid?",[])
        self.assertIn("check ALL relevant references",prompt)
        self.assertIn("Do not confuse a prerequisite",prompt)
        self.assertIn("Never invent a formula",prompt)


class V4LegacyPDFSafetyTests(unittest.TestCase):
    def test_symbol_font_mapping_is_font_specific(self):
        from knowledge_service.document_parser import _decoded_pdf_chars
        sample = [
            {"fontname": "ABCD+SymbolMT", "text": "\uf077"},
            {"fontname": "ABCD+SymbolMT", "text": "\uf061"},
            {"fontname": "ABCD+SymbolMT", "text": "\uf03d"},
            {"fontname": "Wingdings-Regular", "text": "\uf06e"},
            {"fontname": "OtherFont", "text": "\uf077"},
        ]
        mapped = _decoded_pdf_chars(sample)
        self.assertEqual([item["text"] for item in mapped], ["ω", "α", "=", "•", "\uf077"])

    def test_legacy_pdf_requires_layout_aware_extractor(self):
        requirements = (ROOT / "requirements.txt").read_text(encoding="utf-8")
        self.assertIn("pdfplumber>=", requirements)
        from knowledge_service.document_parser import _extract_pdf_pages
        source = __import__("inspect").getsource(_extract_pdf_pages)
        self.assertIn("cannot", source)
        self.assertIn("_decoded_pdf_chars", source)

    def test_garbled_mcq_is_not_displayable(self):
        from common.quiz_engine import _mcq_option_quality
        bad = [
            "Choose i thatmaximizes )|( xiP \uf077 .ThisyieldstheBayes’decisionrule.",
            "Using \uf034 \uf038 otherwise.Wehave \uf077 \uf061 \uf06c.",
            "The Bayes Decision Rule combines Fig.",
            "Evidence: P(ω | x) = p(x) j j .",
        ]
        for value in bad:
            with self.subTest(value=value):
                self.assertFalse(_mcq_option_quality(value))
        self.assertTrue(_mcq_option_quality("Choose the class with the higher posterior probability."))

    def test_bayes_risk_fact_can_be_used_without_private_symbols(self):
        from common.quiz_engine import _usable_true_false_fact, _tf_topic_match
        source = "One wish to vary the choice α(x) to minimize the Bayes’ risk."
        chunk = {"section": "Bayes decision theory", "text": source, "facts": [{"text": source}]}
        self.assertTrue(_usable_true_false_fact(source))
        self.assertTrue(_tf_topic_match("Bayes risk", source, chunk))
        self.assertFalse(_tf_topic_match("normal density classifier", source, chunk))


class V5ExerciseQualityTests(unittest.TestCase):
    @staticmethod
    def _chunk(text, page=18, section="Vocabulary"):
        return {
            "text": text, "source_file": "Intro_2026.pdf", "section": section,
            "chapter": "Intro", "page": page, "chunk_id": f"chunk-{page}",
            "knowledge_key": f"knowledge-{page}", "relevance": 0.9,
            "facts": [{"text": text, "fact_key": f"fact-{page}"}],
        }

    def test_point_source_uses_two_independent_key_points_and_partial_credit(self):
        from common.quiz_engine import _prepare_short_answer_quiz
        fact = "A point source is infinitely small and emits light evenly in all direction."
        chunk = self._chunk(fact)
        focus = FocusFact(fact, "f", chunk)
        quiz, citations = _prepare_short_answer_quiz({
            "question": "According to the course vocabulary, what are the defining characteristics of a point source?"
        }, focus, [chunk], "point source")
        self.assertEqual(len(quiz["key_points"]), 2)
        self.assertEqual(len(citations), 1)
        self.assertIn("infinitely small", quiz["key_points"][0])
        self.assertIn("emits light", quiz["key_points"][1])
        validate_grounding(quiz, [chunk], [], focus)
        result = local_semantic_grade("It is infinitely small.", quiz["reference"], 1, quiz["key_points"])
        self.assertEqual(result["score"], 50)
        self.assertEqual([p["credit"] for p in result["key_point_results"]], [100, 0])
        raw = json.dumps({
            "is_correct": True, "score": 100, "reasoning": "The first property is covered. The second property is absent.",
            "correct_answer": quiz["reference"],
            "key_point_results": [
                {"index": 1, "credit": 100, "comment": "The first property is covered."},
                {"index": 2, "credit": 0, "comment": "The second property is absent."},
            ], "citations": [1], "insufficient_evidence": False,
        })
        normalized = normalize_grading(raw, quiz["reference"], 1, quiz["key_points"], "It is infinitely small.")
        self.assertEqual(normalized["score"], 50)

    def test_lamp_definition_stays_on_topic_even_after_used_by_mcq(self):
        fact = (
            "In building services, the lamp is the electrical component that produce light; "
            "the luminaire is the enclosure that holds and protects the lamp."
        )
        unrelated = "A point source is infinitely small and emits light evenly in all direction."
        chunk = self._chunk(fact)
        chunk["facts"].append({"text": unrelated, "fact_key": "point-source"})
        chunk["text"] = fact + " " + unrelated
        wrong_agent = StaticRouter(AgentKind.QUIZ, json.dumps({
            "type": "short_answer", "question": "What are the defining characteristics of a point source?",
            "key_points": [unrelated], "reference": unrelated, "explanation": unrelated,
        }))
        with patch.dict(os.environ, {"COZE_QUIZ_TYPES": "short_answer"}):
            got = generate_quiz(
                "short_answer", "definition of lamp", {"results": [chunk]}, [],
                {"facts": ["fact-18"]}, wrong_agent, "student", generation_seed="used-before",
            )
        self.assertIn("lamp", got.quiz["question"].casefold())
        self.assertNotIn("point source", got.quiz["question"].casefold())
        self.assertEqual(len(got.quiz["key_points"]), 1)
        self.assertIn("electrical component", got.quiz["reference"])
        validate_grounding(got.quiz, [chunk], [], got.focus)

    def test_fill_blank_rejects_outline_and_grammar_candidates(self):
        from common.quiz_engine import _fill_source_quality, _instructional_fill_answer, _fill_blank_answer
        outline = "Making decisions based on minimum risks 6."
        self.assertFalse(_fill_source_quality(outline))
        self.assertFalse(_instructional_fill_answer("Making"))
        self.assertFalse(_instructional_fill_answer("possible"))
        valid = "For some lamp types, it is possible to change the colour properties by using fluorescent coatings of different chemical composition."
        self.assertEqual(_fill_blank_answer(FocusFact(valid, "f", self._chunk(valid)), [], "lamps"), "fluorescent coatings")
        wrong = {"type":"fill_blank", "question":"_____ decisions based on minimum risks 6.",
                 "answer":"Making", "explanation":outline}
        with self.assertRaises(QuizValidationError):
            validate_grounding(wrong,[self._chunk(outline)],[],FocusFact(outline,"f",self._chunk(outline)))

    def test_fill_blank_definition_is_a_technical_concept(self):
        fact = (
            "In building services, the lamp is the electrical component that produce light; "
            "the luminaire is the enclosure that holds and protects the lamp."
        )
        chunk = self._chunk(fact)
        got = generate_quiz("fill_blank", "definition of lamp", {"results":[chunk]}, [], {},
                            DisabledRouter(), "student", allow_agent=False)
        self.assertEqual(got.quiz["answer"], "luminaire")
        validate_grounding(got.quiz, [chunk], [], got.focus)

    def test_bayes_matching_collects_supported_neighbouring_page_definitions(self):
        from common.quiz_engine import _matching_neighborhood_pairs
        risk = self._chunk(
            "Sounds reasonable to take weighted average of the loss based on the posterior probability for a given x. "
            "This weighted average/conditional expected value is the conditional risk.", page=17,
            section="Bayes Decision Theory",
        )
        risk["source_file"] = "Bayes.pdf"
        risk["facts"] = [{"text":"This weighted average/conditional expected value is the conditional risk.","fact_key":"risk"}]
        error = self._chunk(
            "If the loss-function is chosen as the error-rate (0-1 loss), it is called minimum error-rate classifier.",
            page=18, section="Bayes Decision Theory")
        error["source_file"] = "Bayes.pdf"
        loss = self._chunk("A loss function states the penalty of choosing a wrong class.",
                           page=12, section="Bayes Decision Theory")
        loss["source_file"] = "Bayes.pdf"
        chunks = [risk, error, loss]
        focus = FocusFact("This weighted average/conditional expected value is the conditional risk.","risk",risk)
        self.assertGreaterEqual(len(_matching_neighborhood_pairs(chunks,focus,"Bayes risk")),3)
        generated = local_quiz("matching",focus,chunks,[],"Bayes risk")[0]
        self.assertEqual(len(generated["left"]),3)
        self.assertEqual(set(generated["left"]), {"conditional risk","minimum error-rate classifier","loss function"})
        validate_grounding(generated,chunks,[],focus)

    def test_existing_labeled_gps_matching_is_still_allowed(self):
        chunk = self._chunk("GPS Satellites send timing signals. GPS components include GPS satellites, control segment, and user receiver.")
        chunk["source_file"] = "Ch3.pdf"
        terms = [
            {"term":t,"definition":d,"category":"GPS Components","source_file":"Ch3.pdf","page":p}
            for t,d,p in [
                ("GPS Satellites", "satellites broadcasting signals used for navigation",72),
                ("Control segment", "ground monitoring stations controlling the satellites",74),
                ("User (GPS Receiver)", "a GPS receiver measuring the arriving satellite signals",77),
            ]
        ]
        quiz, cit = local_quiz("matching",FocusFact(chunk["text"],"f",chunk),[chunk],terms,"GPS")
        self.assertEqual(len(quiz["left"]),3)
        self.assertEqual(len(cit),3)
        validate_grounding(quiz,[chunk],terms,FocusFact(chunk["text"],"f",chunk))

    def test_default_retrieval_does_not_add_neighbor_context(self):
        # Repository API compatibility check: new option defaults to zero.
        import inspect
        from knowledge_service.repository import KnowledgeRepository
        from common.knowledge_client import KnowledgeClient
        self.assertEqual(inspect.signature(KnowledgeRepository.search).parameters["neighbor_pages"].default,0)
        self.assertEqual(inspect.signature(KnowledgeClient.search).parameters["neighbor_pages"].default,0)


class WindowsStartupContractTests(unittest.TestCase):
    def test_direct_app_entrypoints_bootstrap_project_root_and_run(self):
        for relative in ("knowledge_service/app.py", "teacher_app/app.py", "student_app/app.py"):
            source = (ROOT / relative).read_text(encoding="utf-8")
            self.assertIn("PROJECT_ROOT = Path(__file__).resolve().parents[1]", source, relative)
            self.assertIn("sys.path.insert(0, str(PROJECT_ROOT))", source, relative)
            self.assertIn('if __name__ == "__main__":', source, relative)
            self.assertIn("application.run(", source, relative)

    def test_windows_launchers_show_failures_and_use_shared_python_resolver(self):
        resolver = (ROOT / "scripts" / "windows_python.bat").read_text(encoding="utf-8")
        self.assertIn("D:\\Python3.11.9\\python.exe", resolver)
        self.assertIn("PYTHONPATH", resolver)
        for relative in ("RUN_KNOWLEDGE.bat", "RUN_TEACHER.bat", "RUN_STUDENT.bat", "START_ALL.bat"):
            source = (ROOT / relative).read_text(encoding="utf-8")
            self.assertIn("windows_", source, relative)
        self.assertIn("pause", (ROOT / "RUN_TEACHER.bat").read_text(encoding="utf-8").lower())
        self.assertIn("pause", (ROOT / "RUN_STUDENT.bat").read_text(encoding="utf-8").lower())


if __name__ == "__main__":
    unittest.main()
