from __future__ import annotations

import json
import unittest
from pathlib import Path
from types import SimpleNamespace

from openpyxl import load_workbook

from common.feedback_builder import build_detailed_feedback, build_record_review, build_review_payload
from teacher_app.kahoot_export import build_kahoot_xlsx, eligible_questions


ROOT = Path(__file__).resolve().parents[1]
CITATIONS = [
    {
        "source_file": "lecture.pdf",
        "section": "GPS Components",
        "page": 12,
        "excerpt": "GPS uses satellites, a control segment and a user receiver.",
    }
]


class DetailedFeedbackTests(unittest.TestCase):
    def assert_common(self, feedback: str, question: str) -> None:
        self.assertIn("[KBot Detailed Feedback v1]", feedback)
        self.assertIn(f"Question: {question}", feedback)
        self.assertIn("Explanation:", feedback)
        self.assertIn("lecture.pdf", feedback)
        self.assertIn("page: 12", feedback)
        self.assertIn("Grading source:", feedback)

    def test_true_false_feedback_contains_both_options_and_correct_option(self):
        question = "GPS is a satellite navigation system."
        feedback = build_detailed_feedback(
            {
                "type": "true_false",
                "question": question,
                "options": ["True", "False"],
                "correct_index": 0,
                "explanation": "The lecture identifies GPS as satellite navigation.",
            },
            "1",
            {
                "is_correct": False,
                "score": 0,
                "reasoning": "The selected option is not supported.",
                "correct_answer": "True",
                "grading_source": "local_objective",
            },
            CITATIONS,
        )
        self.assert_common(feedback, question)
        self.assertIn("A. True [CORRECT]", feedback)
        self.assertIn("B. False [STUDENT SELECTED]", feedback)
        self.assertIn("Correct option: A. True", feedback)
        self.assertIn("Other options: B. False", feedback)

    def test_mcq_feedback_contains_all_options_and_recovery_markers(self):
        question = "Which component provides ranging signals?"
        feedback = build_detailed_feedback(
            {
                "type": "mcq",
                "question": question,
                "options": ["Receiver", "Satellite", "Database", "Router"],
                "correct_index": 1,
                "explanation": "Satellites transmit the ranging signals used by receivers.",
            },
            0,
            {
                "is_correct": False,
                "score": 0,
                "reasoning": "The receiver uses the signal but does not originate it.",
                "correct_answer": "Satellite",
                "grading_source": "local_objective",
            },
            CITATIONS,
        )
        self.assert_common(feedback, question)
        for text in ("A. Receiver", "B. Satellite", "C. Database", "D. Router"):
            self.assertIn(text, feedback)
        self.assertIn("B. Satellite [CORRECT]", feedback)
        self.assertIn("A. Receiver [STUDENT SELECTED]", feedback)
        self.assertIn("Other options: A. Receiver | C. Database | D. Router", feedback)

    def test_matching_feedback_contains_submitted_and_correct_pairs(self):
        question = "Match each GPS segment to its role."
        feedback = build_detailed_feedback(
            {
                "type": "matching",
                "question": question,
                "left": ["Space", "Control"],
                "right": ["Monitors satellites", "Transmits signals"],
                "correct_mapping": [1, 0],
                "explanation": "The space segment transmits signals; control monitors the constellation.",
            },
            [0, 1],
            {
                "is_correct": False,
                "score": 0,
                "reasoning": "0 of 2 pairs match the course-note definitions.",
                "correct_answer": ["Space -> Transmits signals", "Control -> Monitors satellites"],
                "grading_source": "local_objective",
            },
            CITATIONS,
        )
        self.assert_common(feedback, question)
        self.assertIn("Submitted pairs:", feedback)
        self.assertIn("Correct pairs:", feedback)
        self.assertIn("Space -> Transmits signals", feedback)
        self.assertIn("Control -> Monitors satellites", feedback)

    def test_fill_blank_feedback_contains_exact_answer(self):
        question = "GPS uses the _____ segment to monitor satellites."
        feedback = build_detailed_feedback(
            {
                "type": "fill_blank",
                "question": question,
                "answer": "control",
                "explanation": "The control segment monitors satellite status.",
            },
            "space",
            {
                "is_correct": False,
                "score": 0,
                "reasoning": "The blank does not match the saved course-note term.",
                "correct_answer": "control",
                "grading_source": "local_objective",
            },
            CITATIONS,
        )
        self.assert_common(feedback, question)
        self.assertIn("Correct answer: control", feedback)

    def test_short_answer_feedback_contains_reference_answer(self):
        question = "Explain the purpose of the GPS control segment."
        reference = "It monitors and maintains the satellite constellation."
        feedback = build_detailed_feedback(
            {
                "type": "short_answer",
                "question": question,
                "reference": reference,
                "explanation": "The control segment checks satellite health and updates navigation data.",
            },
            "It controls GPS.",
            {
                "is_correct": False,
                "score": 40,
                "reasoning": "The answer is too broad and omits monitoring and maintenance.",
                "correct_answer": reference,
                "grading_source": "coze_grading_agent",
            },
            CITATIONS,
        )
        self.assert_common(feedback, question)
        self.assertIn(f"Reference answer: {reference}", feedback)


class StudentReviewTests(unittest.TestCase):
    def test_mcq_review_contains_question_all_options_selection_and_correct_option(self):
        quiz = {
            "type": "mcq",
            "question": "Which item is listed in the notes?",
            "options": ["Alpha", "Beta", "Gamma", "Delta"],
            "correct_index": 3,
        }
        review = build_review_payload(
            quiz,
            "3",
            {"is_correct": True, "score": 100, "correct_answer": "Delta"},
        )
        self.assertEqual(review["question_type"], "MCQ")
        self.assertEqual(review["question"], quiz["question"])
        self.assertEqual(review["student_answer"], "D. Delta")
        self.assertEqual(len(review["options"]), 4)
        self.assertTrue(review["options"][3]["is_selected"])
        self.assertTrue(review["options"][3]["is_correct"])
        self.assertEqual(review["correct_answer"], "D. Delta")
        self.assertNotIn("generation_source", review)
        self.assertNotIn("grading_source", review)
        self.assertNotIn("citations", review)

    def test_persisted_history_review_does_not_expose_internal_sources(self):
        record = SimpleNamespace(
            question_type="true_false",
            question="GPS uses satellites.",
            answer="0",
            options_or_pairs=json.dumps({"options": ["True", "False"], "correct_index": 0}),
            correct_answer="True",
            is_correct=True,
            score=100,
            teacher_feedback="Well done.",
            generation_source="local_fallback",
            grading_source="local_objective",
            ai_feedback="full teacher audit",
        )
        review = build_record_review(record)
        self.assertEqual(review["student_answer"], "A. True")
        self.assertEqual(review["correct_answer"], "A. True")
        self.assertEqual(review["teacher_feedback"], "Well done.")
        self.assertNotIn("generation_source", review)
        self.assertNotIn("grading_source", review)
        self.assertNotIn("ai_feedback", review)

    def test_persisted_mcq_review_recovers_options_selection_and_correct_option(self):
        record = SimpleNamespace(
            question_type="MCQ",
            question="Which layer provides routing?",
            answer="1",
            options_or_pairs=json.dumps({
                "options": ["Application", "Network", "Transport", "Physical"],
                "correct_index": 1,
            }),
            correct_answer="Network",
            is_correct=True,
            score=100,
            teacher_feedback="",
        )
        review = build_record_review(record)
        self.assertEqual(review["question_type"], "MCQ")
        self.assertEqual(review["student_answer"], "B. Network")
        self.assertEqual(review["correct_answer"], "B. Network")
        self.assertEqual([option["label"] for option in review["options"]], ["A", "B", "C", "D"])
        self.assertTrue(review["options"][1]["is_selected"])
        self.assertTrue(review["options"][1]["is_correct"])

    def test_persisted_matching_review_recovers_submitted_and_correct_pairs(self):
        record = SimpleNamespace(
            question_type="Matching",
            question="Match the GPS components.",
            answer=json.dumps([1, 0]),
            options_or_pairs=json.dumps({
                "left": ["Control segment", "Space segment"],
                "right": ["Satellites", "Monitoring stations"],
                "correct_mapping": [1, 0],
            }),
            correct_answer="",
            is_correct=True,
            score=100,
            teacher_feedback="",
        )
        review = build_record_review(record)
        self.assertEqual(review["question_type"], "Matching")
        self.assertEqual(review["submitted_pairs"], [
            "Control segment -> Monitoring stations",
            "Space segment -> Satellites",
        ])
        self.assertEqual(review["correct_pairs"], [
            "Control segment -> Monitoring stations",
            "Space segment -> Satellites",
        ])
        self.assertEqual(review["student_answer"], "Control segment -> Monitoring stations; Space segment -> Satellites")


class KahootRecoveryTests(unittest.TestCase):
    def test_mcq_can_be_exported_from_detailed_feedback_when_structured_options_are_missing(self):
        quiz = {
            "type": "mcq",
            "question": "Which component transmits GPS signals?",
            "options": ["Control station", "Satellite", "User database", "Map server"],
            "correct_index": 1,
            "explanation": "The satellite broadcasts the ranging signal.",
        }
        feedback = build_detailed_feedback(
            quiz,
            0,
            {
                "is_correct": False,
                "score": 0,
                "reasoning": "The selected option is not correct.",
                "correct_answer": "Satellite",
                "grading_source": "local_objective",
            },
            CITATIONS,
        )
        record = SimpleNamespace(
            question_type="mcq",
            question=quiz["question"],
            options_or_pairs="",
            correct_answer="",
            ai_feedback=feedback,
        )
        questions = eligible_questions([record])
        self.assertEqual(len(questions), 1)
        self.assertEqual(questions[0]["options"], quiz["options"])
        self.assertEqual(questions[0]["correct"], "2")

        output, count = build_kahoot_xlsx([record], ROOT / "teacher_app" / "kahoot_template.xlsx")
        self.assertEqual(count, 1)
        workbook = load_workbook(output)
        sheet = workbook.active
        self.assertEqual(sheet.cell(9, 1).value, quiz["question"])
        self.assertEqual(sheet.cell(9, 3).value, "Satellite")
        self.assertEqual(str(sheet.cell(9, 7).value), "2")

    def test_true_false_can_recover_correct_option_from_detailed_feedback(self):
        quiz = {
            "type": "true_false",
            "question": "GPS uses satellites.",
            "options": ["True", "False"],
            "correct_index": 0,
            "explanation": "The lecture describes GPS as satellite navigation.",
        }
        feedback = build_detailed_feedback(
            quiz,
            1,
            {
                "is_correct": False,
                "score": 0,
                "reasoning": "The selected option is not supported.",
                "correct_answer": "True",
                "grading_source": "local_objective",
            },
            CITATIONS,
        )
        record = SimpleNamespace(
            question_type="true_false",
            question=quiz["question"],
            options_or_pairs="",
            correct_answer="",
            ai_feedback=feedback,
        )
        recovered = eligible_questions([record])[0]
        self.assertEqual(recovered["options"][:2], ["True", "False"])
        self.assertEqual(recovered["correct"], "1")

    def test_feedback_recovery_accepts_combined_correct_and_selected_tag(self):
        quiz = {
            "type": "mcq",
            "question": "Which answer is correct?",
            "options": ["First", "Second", "Third", "Fourth"],
            "correct_index": 2,
            "explanation": "Third is the supported answer.",
        }
        feedback = build_detailed_feedback(
            quiz,
            2,
            {
                "is_correct": True,
                "score": 100,
                "reasoning": "Correct.",
                "correct_answer": "Third",
                "grading_source": "local_objective",
            },
            CITATIONS,
        )
        self.assertIn("C. Third [CORRECT; STUDENT SELECTED]", feedback)
        record = SimpleNamespace(
            question_type="mcq",
            question=quiz["question"],
            options_or_pairs="",
            correct_answer="",
            ai_feedback=feedback,
        )
        recovered = eligible_questions([record])[0]
        self.assertEqual(recovered["options"], quiz["options"])
        self.assertEqual(recovered["correct"], "3")

    def test_structured_fields_remain_the_primary_export_source(self):
        record = SimpleNamespace(
            question_type="mcq",
            question="Which statement is correct?",
            options_or_pairs=json.dumps({"options": ["One", "Two", "Three", "Four"]}),
            correct_answer="Three",
            ai_feedback="unrelated old feedback",
        )
        question = eligible_questions([record])[0]
        self.assertEqual(question["options"], ["One", "Two", "Three", "Four"])
        self.assertEqual(question["correct"], "3")

    def test_structured_correct_index_is_recovered_when_correct_answer_text_is_missing(self):
        record = SimpleNamespace(
            question_type="mcq",
            question="Which option is correct?",
            options_or_pairs=json.dumps({
                "options": ["Alpha", "Beta", "Gamma", "Delta"],
                "correct_index": 0,
            }),
            correct_answer="",
            ai_feedback="",
        )
        question = eligible_questions([record])[0]
        self.assertEqual(question["correct"], "1")


    def test_duplicate_records_are_deduplicated_by_normalized_question_content(self):
        records = [
            SimpleNamespace(
                question_type="mcq",
                question="Which option is correct?",
                options_or_pairs=json.dumps({
                    "options": ["Alpha", "Beta", "Gamma", "Delta"],
                    "correct_index": 1,
                }),
                correct_answer="",
                ai_feedback="",
            ),
            SimpleNamespace(
                question_type="MCQ",
                question="  Which   OPTION is correct?  ",
                options_or_pairs=json.dumps({
                    "options": [" alpha ", "BETA", "Gamma", "Delta"],
                    "correct_index": 1,
                }),
                correct_answer="",
                ai_feedback="",
            ),
        ]
        questions = eligible_questions(records)
        self.assertEqual(len(questions), 1)
        self.assertEqual(questions[0]["correct"], "2")

    def test_same_stem_with_different_options_is_not_deduplicated(self):
        records = [
            SimpleNamespace(
                question_type="mcq",
                question="Which option is correct?",
                options_or_pairs=json.dumps({
                    "options": ["Alpha", "Beta", "Gamma", "Delta"],
                    "correct_index": 1,
                }),
                correct_answer="",
                ai_feedback="",
            ),
            SimpleNamespace(
                question_type="mcq",
                question="Which option is correct?",
                options_or_pairs=json.dumps({
                    "options": ["Alpha", "Second", "Gamma", "Delta"],
                    "correct_index": 1,
                }),
                correct_answer="",
                ai_feedback="",
            ),
        ]
        questions = eligible_questions(records)
        self.assertEqual(len(questions), 2)


    def test_long_correct_mcq_option_uses_saved_correct_index_before_xlsx_truncation(self):
        long_correct = "Some primitive LBS services use a single base station, with a radius of inaccuracy, to determine a phone's location"
        options = [
            "Cell identity only",
            "Satellite orbit prediction",
            "Network database lookup",
            long_correct,
        ]
        record = SimpleNamespace(
            question_type="mcq",
            question='Which item is listed under "Control plane locating" in the course notes?',
            options_or_pairs=json.dumps({"options": options, "correct_index": 3}),
            correct_answer=long_correct,
            ai_feedback="",
        )
        questions = eligible_questions([record])
        self.assertEqual(len(questions), 1)
        self.assertEqual(questions[0]["correct"], "4")
        self.assertEqual(len(questions[0]["options"][3]), 75)

        output, count = build_kahoot_xlsx([record], ROOT / "teacher_app" / "kahoot_template.xlsx")
        self.assertEqual(count, 1)
        workbook = load_workbook(output)
        sheet = workbook.active
        self.assertEqual(str(sheet.cell(9, 7).value), "4")
        self.assertEqual(sheet.cell(9, 5).value, long_correct[:75])

    def test_non_kahoot_question_types_are_not_forced_into_quiz_template(self):
        records = [
            SimpleNamespace(
                question_type=kind,
                question="Question",
                options_or_pairs="",
                correct_answer="Answer",
                ai_feedback="[KBot Detailed Feedback v1]",
            )
            for kind in ("matching", "fill_blank", "short_answer")
        ]
        self.assertEqual(eligible_questions(records), [])


if __name__ == "__main__":
    unittest.main()
